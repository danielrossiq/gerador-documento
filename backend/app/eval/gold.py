"""Conjunto ouro da IA integrada (plano ia-integrada-plano.md v0.4, item 1): 15 modelos FICTÍCIOS anotados.

Cada documento é escrito como modelos de parágrafo com marcadores `{campo}`. O texto final troca cada marcador pelo
valor de exemplo, e cada troca é uma ocorrência anotada: bloco, posição e trecho exatos. Assim a anotação nunca
diverge do texto. Os campos dizem rótulo, tipo (vocabulário FIELD_TYPES), parte (contratante, locador…) e grupo do
formulário. Nomes, documentos, endereços e valores são inventados; CPF e CNPJ são gerados com dígitos válidos a
partir de números fixos, sem corresponder a pessoas reais conhecidas.

Os .docx são montados em memória com python-docx (determinístico), então nada binário vai para o repositório.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

from docx import Document

from app.modules.fields.types import FIELD_TYPES

CATEGORIES = ("serviços", "locação", "procuração", "declaração", "recibo", "termo", "proposta", "RH")
_MARK = re.compile(r"\{([a-z0-9_]+)\}")


def cpf(base9: str) -> str:
    """CPF com dígitos verificadores válidos a partir de 9 dígitos fixos (fictício)."""
    d = [int(c) for c in base9]
    for n in (10, 11):
        s = sum(v * w for v, w in zip(d, range(n, 1, -1), strict=False)) % 11
        d.append(0 if s < 2 else 11 - s)
    t = "".join(map(str, d))
    return f"{t[:3]}.{t[3:6]}.{t[6:9]}-{t[9:]}"


def cnpj(base12: str) -> str:
    """CNPJ com dígitos verificadores válidos a partir de 12 dígitos fixos (fictício)."""
    d = [int(c) for c in base12]
    for weights in ([5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2], [6, 5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]):
        s = sum(v * w for v, w in zip(d, weights, strict=True)) % 11
        d.append(0 if s < 2 else 11 - s)
    t = "".join(map(str, d))
    return f"{t[:2]}.{t[2:5]}.{t[5:8]}/{t[8:12]}-{t[12:]}"


@dataclass(frozen=True)
class GoldField:
    name: str
    label: str
    type: str
    value: str
    party: str | None = None  # papel da parte (contratante, locador…), quando o campo pertence a uma
    group: str = "Dados do documento"


@dataclass(frozen=True)
class Occurrence:
    field: str
    block: str
    start: int
    end: int
    text: str


@dataclass
class GoldDoc:
    id: str
    category: str
    title: str
    paragraphs: list[str]
    fields: list[GoldField]
    header: str | None = None
    occurrences: list[Occurrence] = field(default_factory=list, init=False)

    def render(self) -> list[tuple[str, str]]:
        """(chave do bloco, texto) na mesma convenção do parser (body/pN; header1/p1). Preenche as ocorrências."""
        values = {f.name: f.value for f in self.fields}
        blocks: list[tuple[str, str]] = []
        self.occurrences = []
        items = ([("header1/p1", self.header)] if self.header else []) + [
            (f"body/p{i}", p) for i, p in enumerate([self.title, *self.paragraphs], 1)
        ]
        for key, template in items:
            out, pos = "", 0
            for m in _MARK.finditer(template):
                out += template[pos : m.start()]
                value = values[m.group(1)]
                self.occurrences.append(Occurrence(m.group(1), key, len(out), len(out) + len(value), value))
                out += value
                pos = m.end()
            blocks.append((key, out + template[pos:]))
        return blocks

    def docx(self) -> bytes:
        doc = Document()
        blocks = dict(self.render())
        if self.header:
            doc.sections[0].header.paragraphs[0].text = blocks["header1/p1"]
        for i in range(1, len(self.paragraphs) + 2):
            doc.add_paragraph(blocks[f"body/p{i}"])
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()


F = GoldField


def _docs() -> list[GoldDoc]:
    return [
        GoldDoc(
            "servicos-ti",
            "serviços",
            "CONTRATO DE PRESTAÇÃO DE SERVIÇOS",
            [
                "CONTRATANTE: {nome_contratante}, inscrita no CPF sob o nº {cpf_contratante}, residente na {endereco_contratante}, CEP {cep_contratante}.",
                "CONTRATADA: {nome_contratada}, inscrita no CNPJ sob o nº {cnpj_contratada}.",
                "CLÁUSULA 1ª – DO OBJETO. A CONTRATADA prestará serviços de {servico} em favor de {nome_contratante}.",
                "CLÁUSULA 2ª – DO VALOR. O valor mensal é de {valor_mensal}, pago até o dia {dia_pagamento} de cada mês.",
                "CLÁUSULA 3ª – DA VIGÊNCIA. O contrato inicia em {data_inicio} e vigora por {prazo_meses} meses.",
                "{cidade}, {data_assinatura}.",
                "{nome_contratante} — CONTRATANTE",
            ],
            [
                F(
                    "nome_contratante",
                    "Nome da contratante",
                    "text_short",
                    "Marina Duarte Campos",
                    "contratante",
                    "Contratante",
                ),
                F(
                    "cpf_contratante",
                    "CPF da contratante",
                    "cpf",
                    cpf("418226305"),
                    "contratante",
                    "Contratante",
                ),
                F(
                    "endereco_contratante",
                    "Endereço da contratante",
                    "text_short",
                    "Rua das Acácias, 120, apto 31",
                    "contratante",
                    "Contratante",
                ),
                F("cep_contratante", "CEP da contratante", "cep", "04512-030", "contratante", "Contratante"),
                F(
                    "nome_contratada",
                    "Nome da contratada",
                    "text_short",
                    "Lumen Sistemas Ltda.",
                    "contratada",
                    "Contratada",
                ),
                F(
                    "cnpj_contratada",
                    "CNPJ da contratada",
                    "cnpj",
                    cnpj("318540920001"),
                    "contratada",
                    "Contratada",
                ),
                F("servico", "Serviço", "text_short", "manutenção de sistemas internos", None, "Objeto"),
                F("valor_mensal", "Valor mensal", "currency", "R$ 3.200,00", None, "Pagamento"),
                F("dia_pagamento", "Dia do pagamento", "number", "10", None, "Pagamento"),
                F("data_inicio", "Data de início", "date", "01/11/2026", None, "Vigência"),
                F("prazo_meses", "Prazo em meses", "number", "12", None, "Vigência"),
                F("cidade", "Cidade", "text_short", "Campinas", None, "Assinatura"),
                F("data_assinatura", "Data da assinatura", "date", "20/10/2026", None, "Assinatura"),
            ],
            header="Lumen Sistemas Ltda. — Contrato de Serviços",
        ),
        GoldDoc(
            "servicos-limpeza",
            "serviços",
            "CONTRATO DE PRESTAÇÃO DE SERVIÇOS DE LIMPEZA",
            [
                "Pelo presente instrumento, {nome_contratante}, CNPJ {cnpj_contratante}, doravante CONTRATANTE, e {nome_contratado}, CPF {cpf_contratado}, doravante CONTRATADO, ajustam o seguinte.",
                "O CONTRATADO realizará a limpeza do imóvel situado na {endereco_imovel} às {dia_semana}s.",
                "Pelos serviços, a CONTRATANTE pagará {valor} por visita, mediante PIX para a chave {chave_pix}.",
                "Contato do CONTRATADO: {telefone_contratado}; e-mail {email_contratado}.",
                "{cidade}, {data_assinatura}.",
            ],
            [
                F(
                    "nome_contratante",
                    "Nome da contratante",
                    "text_short",
                    "Clínica Bem Estar Ltda.",
                    "contratante",
                    "Contratante",
                ),
                F(
                    "cnpj_contratante",
                    "CNPJ da contratante",
                    "cnpj",
                    cnpj("472091330001"),
                    "contratante",
                    "Contratante",
                ),
                F(
                    "nome_contratado",
                    "Nome do contratado",
                    "text_short",
                    "Rafael Moreira Lins",
                    "contratado",
                    "Contratado",
                ),
                F("cpf_contratado", "CPF do contratado", "cpf", cpf("205917348"), "contratado", "Contratado"),
                F(
                    "endereco_imovel",
                    "Endereço do imóvel",
                    "text_short",
                    "Avenida Brasil, 2.450, sala 4",
                    None,
                    "Objeto",
                ),
                F("dia_semana", "Dia da semana", "select", "segunda-feira", None, "Objeto"),
                F("valor", "Valor por visita", "currency", "R$ 180,00", None, "Pagamento"),
                F(
                    "chave_pix",
                    "Chave PIX",
                    "text_short",
                    "rafael.lins@exemplo.test",
                    "contratado",
                    "Pagamento",
                ),
                F(
                    "telefone_contratado",
                    "Telefone do contratado",
                    "phone",
                    "(19) 98765-4321",
                    "contratado",
                    "Contratado",
                ),
                F(
                    "email_contratado",
                    "E-mail do contratado",
                    "email",
                    "rafael.lins@exemplo.test",
                    "contratado",
                    "Contratado",
                ),
                F("cidade", "Cidade", "text_short", "Sorocaba", None, "Assinatura"),
                F("data_assinatura", "Data da assinatura", "date", "18/10/2026", None, "Assinatura"),
            ],
        ),
        GoldDoc(
            "locacao-residencial",
            "locação",
            "CONTRATO DE LOCAÇÃO RESIDENCIAL",
            [
                "LOCADOR: {nome_locador}, {nacionalidade_locador}, {estado_civil_locador}, CPF {cpf_locador}.",
                "LOCATÁRIA: {nome_locataria}, {nacionalidade_locataria}, {estado_civil_locataria}, CPF {cpf_locataria}.",
                "OBJETO: o imóvel residencial situado na {endereco_imovel}, CEP {cep_imovel}.",
                "PRAZO: {prazo_meses} meses, a partir de {data_inicio}.",
                "ALUGUEL: {valor_aluguel} mensais, reajustado anualmente pelo {indice_reajuste}.",
                "GARANTIA: {garantia}.",
                "{cidade}, {data_assinatura}.",
            ],
            [
                F(
                    "nome_locador",
                    "Nome do locador",
                    "text_short",
                    "Otávio Ribeiro Neves",
                    "locador",
                    "Locador",
                ),
                F(
                    "nacionalidade_locador",
                    "Nacionalidade do locador",
                    "text_short",
                    "brasileiro",
                    "locador",
                    "Locador",
                ),
                F(
                    "estado_civil_locador",
                    "Estado civil do locador",
                    "select",
                    "casado",
                    "locador",
                    "Locador",
                ),
                F("cpf_locador", "CPF do locador", "cpf", cpf("731004286"), "locador", "Locador"),
                F(
                    "nome_locataria",
                    "Nome da locatária",
                    "text_short",
                    "Beatriz Fontes Arantes",
                    "locatária",
                    "Locatária",
                ),
                F(
                    "nacionalidade_locataria",
                    "Nacionalidade da locatária",
                    "text_short",
                    "brasileira",
                    "locatária",
                    "Locatária",
                ),
                F(
                    "estado_civil_locataria",
                    "Estado civil da locatária",
                    "select",
                    "solteira",
                    "locatária",
                    "Locatária",
                ),
                F("cpf_locataria", "CPF da locatária", "cpf", cpf("590382417"), "locatária", "Locatária"),
                F(
                    "endereco_imovel",
                    "Endereço do imóvel",
                    "text_short",
                    "Rua Monte Alegre, 88, casa 2",
                    None,
                    "Imóvel",
                ),
                F("cep_imovel", "CEP do imóvel", "cep", "13010-210", None, "Imóvel"),
                F("prazo_meses", "Prazo em meses", "number", "30", None, "Prazo"),
                F("data_inicio", "Data de início", "date", "05/11/2026", None, "Prazo"),
                F("valor_aluguel", "Valor do aluguel", "currency", "R$ 2.750,00", None, "Aluguel"),
                F("indice_reajuste", "Índice de reajuste", "select", "IPCA", None, "Aluguel"),
                F("garantia", "Garantia", "select", "caução de três aluguéis", None, "Garantia"),
                F("cidade", "Cidade", "text_short", "Campinas", None, "Assinatura"),
                F("data_assinatura", "Data da assinatura", "date", "22/10/2026", None, "Assinatura"),
            ],
        ),
        GoldDoc(
            "locacao-comercial",
            "locação",
            "CONTRATO DE LOCAÇÃO NÃO RESIDENCIAL",
            [
                "LOCADORA: {nome_locadora}, CNPJ {cnpj_locadora}.",
                "LOCATÁRIA: {nome_locataria}, CNPJ {cnpj_locataria}, representada por {representante}, CPF {cpf_representante}.",
                "A LOCATÁRIA usará a sala situada na {endereco_imovel} exclusivamente para {finalidade}.",
                "O aluguel mensal é de {valor_aluguel}, com vencimento no dia {dia_vencimento}. O prazo é de {prazo_meses} meses.",
                "{cidade}, {data_assinatura}.",
            ],
            [
                F(
                    "nome_locadora",
                    "Nome da locadora",
                    "text_short",
                    "Pátio Central Empreendimentos S.A.",
                    "locadora",
                    "Locadora",
                ),
                F("cnpj_locadora", "CNPJ da locadora", "cnpj", cnpj("119284560001"), "locadora", "Locadora"),
                F(
                    "nome_locataria",
                    "Nome da locatária",
                    "text_short",
                    "Estúdio Traço Fino Ltda.",
                    "locatária",
                    "Locatária",
                ),
                F(
                    "cnpj_locataria",
                    "CNPJ da locatária",
                    "cnpj",
                    cnpj("605738210001"),
                    "locatária",
                    "Locatária",
                ),
                F(
                    "representante",
                    "Representante da locatária",
                    "text_short",
                    "Joana Prado Vilela",
                    "locatária",
                    "Locatária",
                ),
                F(
                    "cpf_representante",
                    "CPF da representante",
                    "cpf",
                    cpf("381650924"),
                    "locatária",
                    "Locatária",
                ),
                F(
                    "endereco_imovel",
                    "Endereço do imóvel",
                    "text_short",
                    "Rua XV de Novembro, 300, sala 1204",
                    None,
                    "Imóvel",
                ),
                F("finalidade", "Finalidade", "text_short", "atividades de design gráfico", None, "Imóvel"),
                F("valor_aluguel", "Valor do aluguel", "currency", "R$ 4.900,00", None, "Aluguel"),
                F("dia_vencimento", "Dia do vencimento", "number", "5", None, "Aluguel"),
                F("prazo_meses", "Prazo em meses", "number", "36", None, "Aluguel"),
                F("cidade", "Cidade", "text_short", "Curitiba", None, "Assinatura"),
                F("data_assinatura", "Data da assinatura", "date", "14/10/2026", None, "Assinatura"),
            ],
        ),
        GoldDoc(
            "procuracao-ad-judicia",
            "procuração",
            "PROCURAÇÃO AD JUDICIA",
            [
                "OUTORGANTE: {nome_outorgante}, {nacionalidade}, {profissao}, CPF {cpf_outorgante}, residente na {endereco_outorgante}.",
                "OUTORGADA: {nome_advogada}, advogada inscrita na {oab_advogada}, com escritório na {endereco_escritorio}.",
                "PODERES: os da cláusula ad judicia et extra, para propor {acao} contra {parte_contraria}.",
                "{cidade}, {data_assinatura}.",
                "{nome_outorgante}",
            ],
            [
                F(
                    "nome_outorgante",
                    "Nome do outorgante",
                    "text_short",
                    "Henrique Sales Bastos",
                    "outorgante",
                    "Outorgante",
                ),
                F(
                    "nacionalidade",
                    "Nacionalidade do outorgante",
                    "text_short",
                    "brasileiro",
                    "outorgante",
                    "Outorgante",
                ),
                F(
                    "profissao",
                    "Profissão do outorgante",
                    "text_short",
                    "engenheiro civil",
                    "outorgante",
                    "Outorgante",
                ),
                F("cpf_outorgante", "CPF do outorgante", "cpf", cpf("627391058"), "outorgante", "Outorgante"),
                F(
                    "endereco_outorgante",
                    "Endereço do outorgante",
                    "text_short",
                    "Rua Ipê Roxo, 45",
                    "outorgante",
                    "Outorgante",
                ),
                F(
                    "nome_advogada",
                    "Nome da advogada",
                    "text_short",
                    "Lívia Carvalho Penna",
                    "outorgada",
                    "Outorgada",
                ),
                F("oab_advogada", "OAB da advogada", "oab", "OAB/SP 123.456", "outorgada", "Outorgada"),
                F(
                    "endereco_escritorio",
                    "Endereço do escritório",
                    "text_short",
                    "Avenida Paulista, 1.000, conj. 81",
                    "outorgada",
                    "Outorgada",
                ),
                F("acao", "Ação", "text_short", "ação de cobrança", None, "Poderes"),
                F(
                    "parte_contraria",
                    "Parte contrária",
                    "text_short",
                    "Construtora Horizonte Ltda.",
                    None,
                    "Poderes",
                ),
                F("cidade", "Cidade", "text_short", "São Paulo", None, "Assinatura"),
                F("data_assinatura", "Data da assinatura", "date", "09/10/2026", None, "Assinatura"),
            ],
        ),
        GoldDoc(
            "substabelecimento",
            "procuração",
            "SUBSTABELECIMENTO COM RESERVA DE PODERES",
            [
                "{nome_substabelecente}, {oab_substabelecente}, substabelece, com reserva de iguais poderes, a {nome_substabelecido}, {oab_substabelecido}, os poderes conferidos por {nome_cliente} nos autos do processo nº {numero_processo}.",
                "{cidade}, {data_assinatura}.",
            ],
            [
                F(
                    "nome_substabelecente",
                    "Nome do advogado que substabelece",
                    "text_short",
                    "Caio Mendes Torres",
                    "substabelecente",
                    "Substabelecente",
                ),
                F(
                    "oab_substabelecente",
                    "OAB de quem substabelece",
                    "oab",
                    "OAB/RJ 98.765",
                    "substabelecente",
                    "Substabelecente",
                ),
                F(
                    "nome_substabelecido",
                    "Nome do advogado substabelecido",
                    "text_short",
                    "Paula Reis Andrade",
                    "substabelecido",
                    "Substabelecido",
                ),
                F(
                    "oab_substabelecido",
                    "OAB do substabelecido",
                    "oab",
                    "OAB/RJ 154.302",
                    "substabelecido",
                    "Substabelecido",
                ),
                F("nome_cliente", "Nome do cliente", "text_short", "Sérgio Vidal Costa", None, "Processo"),
                F(
                    "numero_processo",
                    "Número do processo",
                    "text_short",
                    "0801234-56.2026.8.19.0001",
                    None,
                    "Processo",
                ),
                F("cidade", "Cidade", "text_short", "Niterói", None, "Assinatura"),
                F("data_assinatura", "Data da assinatura", "date", "30/09/2026", None, "Assinatura"),
            ],
        ),
        GoldDoc(
            "declaracao-hipossuficiencia",
            "declaração",
            "DECLARAÇÃO DE HIPOSSUFICIÊNCIA",
            [
                "Eu, {nome_declarante}, {estado_civil}, CPF {cpf_declarante}, residente na {endereco}, CEP {cep}, declaro, sob as penas da lei, que não tenho condições de arcar com as custas do processo sem prejuízo do meu sustento.",
                "Minha renda mensal é de {renda_mensal}.",
                "{cidade}, {data_assinatura}.",
                "{nome_declarante}",
            ],
            [
                F(
                    "nome_declarante",
                    "Nome da declarante",
                    "text_short",
                    "Cláudia Ferraz Nogueira",
                    "declarante",
                    "Declarante",
                ),
                F("estado_civil", "Estado civil", "select", "divorciada", "declarante", "Declarante"),
                F("cpf_declarante", "CPF da declarante", "cpf", cpf("903425716"), "declarante", "Declarante"),
                F("endereco", "Endereço", "text_short", "Travessa do Sol, 7", "declarante", "Declarante"),
                F("cep", "CEP", "cep", "60110-090", "declarante", "Declarante"),
                F("renda_mensal", "Renda mensal", "currency", "R$ 1.612,40", "declarante", "Renda"),
                F("cidade", "Cidade", "text_short", "Fortaleza", None, "Assinatura"),
                F("data_assinatura", "Data da assinatura", "date", "11/10/2026", None, "Assinatura"),
            ],
        ),
        GoldDoc(
            "declaracao-residencia",
            "declaração",
            "DECLARAÇÃO DE RESIDÊNCIA",
            [
                "Eu, {nome_declarante}, CPF {cpf_declarante}, telefone {telefone}, declaro que resido desde {data_desde} no endereço {endereco}, {bairro}, {cidade}.",
                "Declaro ainda que {nome_morador}, CPF {cpf_morador}, reside no mesmo endereço.",
                "{cidade}, {data_assinatura}.",
            ],
            [
                F(
                    "nome_declarante",
                    "Nome do declarante",
                    "text_short",
                    "Antônio Valente Rocha",
                    "declarante",
                    "Declarante",
                ),
                F("cpf_declarante", "CPF do declarante", "cpf", cpf("154883260"), "declarante", "Declarante"),
                F("telefone", "Telefone", "phone", "(85) 3244-1020", "declarante", "Declarante"),
                F("data_desde", "Desde", "date", "15/03/2019", None, "Endereço"),
                F("endereco", "Endereço", "text_short", "Rua Padre Valdevino, 1.410", None, "Endereço"),
                F("bairro", "Bairro", "text_short", "Aldeota", None, "Endereço"),
                F("cidade", "Cidade", "text_short", "Fortaleza", None, "Endereço"),
                F(
                    "nome_morador",
                    "Nome do morador",
                    "text_short",
                    "Lúcia Valente Rocha",
                    "morador",
                    "Morador",
                ),
                F("cpf_morador", "CPF do morador", "cpf", cpf("482017395"), "morador", "Morador"),
                F("data_assinatura", "Data da assinatura", "date", "02/10/2026", None, "Assinatura"),
            ],
        ),
        GoldDoc(
            "recibo-honorarios",
            "recibo",
            "RECIBO",
            [
                "Recebi de {nome_pagador}, CPF {cpf_pagador}, a importância de {valor} ({valor_extenso}), referente a {referencia}.",
                "Pagamento realizado em {data_pagamento} por {forma_pagamento}.",
                "{cidade}, {data_assinatura}.",
                "{nome_recebedor} — {oab_recebedor}",
            ],
            [
                F(
                    "nome_pagador",
                    "Nome do pagador",
                    "text_short",
                    "Fernando Albuquerque Sá",
                    "pagador",
                    "Pagador",
                ),
                F("cpf_pagador", "CPF do pagador", "cpf", cpf("316209874"), "pagador", "Pagador"),
                F("valor", "Valor", "currency", "R$ 2.500,00", None, "Pagamento"),
                F(
                    "valor_extenso",
                    "Valor por extenso",
                    "text_short",
                    "dois mil e quinhentos reais",
                    None,
                    "Pagamento",
                ),
                F(
                    "referencia",
                    "Referência",
                    "text_short",
                    "honorários da primeira fase do processo",
                    None,
                    "Pagamento",
                ),
                F("data_pagamento", "Data do pagamento", "date", "16/10/2026", None, "Pagamento"),
                F(
                    "forma_pagamento",
                    "Forma de pagamento",
                    "select",
                    "transferência bancária",
                    None,
                    "Pagamento",
                ),
                F("cidade", "Cidade", "text_short", "Belo Horizonte", None, "Assinatura"),
                F("data_assinatura", "Data da assinatura", "date", "16/10/2026", None, "Assinatura"),
                F(
                    "nome_recebedor",
                    "Nome de quem recebe",
                    "text_short",
                    "Bruno Teixeira Lago",
                    "recebedor",
                    "Recebedor",
                ),
                F("oab_recebedor", "OAB de quem recebe", "oab", "OAB/MG 87.210", "recebedor", "Recebedor"),
            ],
        ),
        GoldDoc(
            "recibo-aluguel",
            "recibo",
            "RECIBO DE ALUGUEL",
            [
                "Recebi de {nome_locatario} a quantia de {valor_aluguel}, referente ao aluguel de {mes_referencia} do imóvel situado na {endereco_imovel}.",
                "{cidade}, {data_assinatura}.",
                "{nome_locador}, CPF {cpf_locador}",
            ],
            [
                F(
                    "nome_locatario",
                    "Nome do locatário",
                    "text_short",
                    "Gustavo Henrique Paes",
                    "locatário",
                    "Locatário",
                ),
                F("valor_aluguel", "Valor do aluguel", "currency", "R$ 1.980,00", None, "Pagamento"),
                F("mes_referencia", "Mês de referência", "text_short", "outubro de 2026", None, "Pagamento"),
                F(
                    "endereco_imovel",
                    "Endereço do imóvel",
                    "text_short",
                    "Rua Dom Pedro II, 512, apto 12",
                    None,
                    "Imóvel",
                ),
                F("cidade", "Cidade", "text_short", "Porto Alegre", None, "Assinatura"),
                F("data_assinatura", "Data da assinatura", "date", "05/10/2026", None, "Assinatura"),
                F(
                    "nome_locador",
                    "Nome do locador",
                    "text_short",
                    "Regina Lopes Sardinha",
                    "locador",
                    "Locador",
                ),
                F("cpf_locador", "CPF do locador", "cpf", cpf("760114283"), "locador", "Locador"),
            ],
        ),
        GoldDoc(
            "termo-confidencialidade",
            "termo",
            "TERMO DE CONFIDENCIALIDADE",
            [
                "{nome_empresa}, CNPJ {cnpj_empresa}, e {nome_parceiro}, CPF {cpf_parceiro}, firmam este termo para proteger as informações trocadas sobre {projeto}.",
                "O sigilo vale por {prazo_anos} anos a partir de {data_inicio}. A multa por descumprimento é de {multa}.",
                "Foro da comarca de {foro}.",
                "{cidade}, {data_assinatura}.",
            ],
            [
                F(
                    "nome_empresa",
                    "Nome da empresa",
                    "text_short",
                    "Nexo Agro Tecnologia Ltda.",
                    "empresa",
                    "Empresa",
                ),
                F("cnpj_empresa", "CNPJ da empresa", "cnpj", cnpj("274503810001"), "empresa", "Empresa"),
                F(
                    "nome_parceiro",
                    "Nome do parceiro",
                    "text_short",
                    "Diego Amaral Prates",
                    "parceiro",
                    "Parceiro",
                ),
                F("cpf_parceiro", "CPF do parceiro", "cpf", cpf("839270461"), "parceiro", "Parceiro"),
                F("projeto", "Projeto", "text_short", "o sistema de irrigação inteligente", None, "Objeto"),
                F("prazo_anos", "Prazo em anos", "number", "5", None, "Prazo"),
                F("data_inicio", "Data de início", "date", "01/12/2026", None, "Prazo"),
                F("multa", "Multa", "currency", "R$ 50.000,00", None, "Penalidade"),
                F("foro", "Foro", "text_short", "Goiânia", None, "Foro"),
                F("cidade", "Cidade", "text_short", "Goiânia", None, "Assinatura"),
                F("data_assinatura", "Data da assinatura", "date", "08/10/2026", None, "Assinatura"),
            ],
        ),
        GoldDoc(
            "termo-acordo",
            "termo",
            "TERMO DE ACORDO E PARCELAMENTO",
            [
                "CREDOR: {nome_credor}, CNPJ {cnpj_credor}. DEVEDOR: {nome_devedor}, CPF {cpf_devedor}.",
                "O DEVEDOR reconhece a dívida de {valor_divida} e pagará em {parcelas} parcelas de {valor_parcela}, vencendo a primeira em {data_primeira}.",
                "O atraso de qualquer parcela gera multa de {multa_percentual} sobre o saldo.",
                "{cidade}, {data_assinatura}.",
            ],
            [
                F(
                    "nome_credor",
                    "Nome do credor",
                    "text_short",
                    "Distribuidora Vale Verde Ltda.",
                    "credor",
                    "Credor",
                ),
                F("cnpj_credor", "CNPJ do credor", "cnpj", cnpj("380162740001"), "credor", "Credor"),
                F(
                    "nome_devedor",
                    "Nome do devedor",
                    "text_short",
                    "Marcos Paulo Gentil",
                    "devedor",
                    "Devedor",
                ),
                F("cpf_devedor", "CPF do devedor", "cpf", cpf("071538924"), "devedor", "Devedor"),
                F("valor_divida", "Valor da dívida", "currency", "R$ 12.600,00", None, "Acordo"),
                F("parcelas", "Número de parcelas", "number", "6", None, "Acordo"),
                F("valor_parcela", "Valor da parcela", "currency", "R$ 2.100,00", None, "Acordo"),
                F("data_primeira", "Vencimento da 1ª parcela", "date", "10/11/2026", None, "Acordo"),
                F("multa_percentual", "Multa por atraso", "text_short", "10%", None, "Acordo"),
                F("cidade", "Cidade", "text_short", "Recife", None, "Assinatura"),
                F("data_assinatura", "Data da assinatura", "date", "19/10/2026", None, "Assinatura"),
            ],
        ),
        GoldDoc(
            "proposta-honorarios",
            "proposta",
            "PROPOSTA DE HONORÁRIOS ADVOCATÍCIOS",
            [
                "À {nome_cliente}, CNPJ {cnpj_cliente}, aos cuidados de {contato_cliente} ({email_contato}).",
                "Apresentamos proposta para {servico_juridico}. Os honorários fixos são de {honorarios_fixos}, mais {percentual_exito} sobre o proveito econômico.",
                "Esta proposta é válida até {data_validade}.",
                "{nome_escritorio} — {nome_advogado}, {oab_advogado}",
            ],
            [
                F(
                    "nome_cliente",
                    "Nome do cliente",
                    "text_short",
                    "Padaria Trigo Dourado Ltda.",
                    "cliente",
                    "Cliente",
                ),
                F("cnpj_cliente", "CNPJ do cliente", "cnpj", cnpj("517820630001"), "cliente", "Cliente"),
                F(
                    "contato_cliente",
                    "Contato do cliente",
                    "text_short",
                    "Sra. Elisa Moura",
                    "cliente",
                    "Cliente",
                ),
                F(
                    "email_contato",
                    "E-mail do contato",
                    "email",
                    "elisa.moura@exemplo.test",
                    "cliente",
                    "Cliente",
                ),
                F(
                    "servico_juridico",
                    "Serviço jurídico",
                    "text_long",
                    "a defesa em reclamação trabalhista em primeira instância",
                    None,
                    "Proposta",
                ),
                F("honorarios_fixos", "Honorários fixos", "currency", "R$ 8.000,00", None, "Honorários"),
                F("percentual_exito", "Percentual de êxito", "text_short", "15%", None, "Honorários"),
                F("data_validade", "Validade", "date", "31/10/2026", None, "Proposta"),
                F(
                    "nome_escritorio",
                    "Nome do escritório",
                    "text_short",
                    "Moura Lins Advogados",
                    "escritório",
                    "Escritório",
                ),
                F(
                    "nome_advogado",
                    "Nome do advogado",
                    "text_short",
                    "Tiago Lins Barreto",
                    "escritório",
                    "Escritório",
                ),
                F("oab_advogado", "OAB do advogado", "oab", "OAB/PE 45.871", "escritório", "Escritório"),
            ],
        ),
        GoldDoc(
            "rh-contrato-trabalho",
            "RH",
            "CONTRATO INDIVIDUAL DE TRABALHO",
            [
                "EMPREGADORA: {nome_empregadora}, CNPJ {cnpj_empregadora}.",
                "EMPREGADA: {nome_empregada}, CPF {cpf_empregada}, residente na {endereco_empregada}.",
                "A EMPREGADA exercerá a função de {cargo}, com jornada de {jornada_horas} horas semanais e salário mensal de {salario}.",
                "Admissão em {data_admissao}, com período de experiência de {experiencia_dias} dias.",
                "{cidade}, {data_assinatura}.",
            ],
            [
                F(
                    "nome_empregadora",
                    "Nome da empregadora",
                    "text_short",
                    "Oficina Criativa Comunicação Ltda.",
                    "empregadora",
                    "Empregadora",
                ),
                F(
                    "cnpj_empregadora",
                    "CNPJ da empregadora",
                    "cnpj",
                    cnpj("623049180001"),
                    "empregadora",
                    "Empregadora",
                ),
                F(
                    "nome_empregada",
                    "Nome da empregada",
                    "text_short",
                    "Sabrina Lacerda Pinto",
                    "empregada",
                    "Empregada",
                ),
                F("cpf_empregada", "CPF da empregada", "cpf", cpf("247806139"), "empregada", "Empregada"),
                F(
                    "endereco_empregada",
                    "Endereço da empregada",
                    "text_short",
                    "Rua Bela Cintra, 980, apto 5",
                    "empregada",
                    "Empregada",
                ),
                F("cargo", "Cargo", "text_short", "assistente administrativa", None, "Contrato"),
                F("jornada_horas", "Jornada semanal", "number", "44", None, "Contrato"),
                F("salario", "Salário", "currency", "R$ 2.840,00", None, "Contrato"),
                F("data_admissao", "Data de admissão", "date", "03/11/2026", None, "Contrato"),
                F("experiencia_dias", "Experiência em dias", "number", "90", None, "Contrato"),
                F("cidade", "Cidade", "text_short", "São Paulo", None, "Assinatura"),
                F("data_assinatura", "Data da assinatura", "date", "27/10/2026", None, "Assinatura"),
            ],
        ),
        GoldDoc(
            "rh-aviso-ferias",
            "RH",
            "AVISO DE FÉRIAS",
            [
                "Prezado(a) {nome_empregado}, matrícula {matricula}, comunicamos que suas férias referentes ao período aquisitivo {periodo_aquisitivo} serão gozadas de {data_inicio_ferias} a {data_fim_ferias}.",
                "O pagamento de {valor_ferias} será feito até {data_pagamento}.",
                "{nome_empresa} — {cidade}, {data_assinatura}.",
            ],
            [
                F(
                    "nome_empregado",
                    "Nome do empregado",
                    "text_short",
                    "Wagner Coelho Diniz",
                    "empregado",
                    "Empregado",
                ),
                F("matricula", "Matrícula", "text_short", "MT-20417", "empregado", "Empregado"),
                F("periodo_aquisitivo", "Período aquisitivo", "text_short", "2025/2026", None, "Férias"),
                F("data_inicio_ferias", "Início das férias", "date", "04/01/2027", None, "Férias"),
                F("data_fim_ferias", "Fim das férias", "date", "02/02/2027", None, "Férias"),
                F("valor_ferias", "Valor das férias", "currency", "R$ 4.210,67", None, "Pagamento"),
                F("data_pagamento", "Data do pagamento", "date", "30/12/2026", None, "Pagamento"),
                F(
                    "nome_empresa",
                    "Nome da empresa",
                    "text_short",
                    "Transportes Rio Claro Ltda.",
                    "empresa",
                    "Empresa",
                ),
                F("cidade", "Cidade", "text_short", "Rio Claro", None, "Assinatura"),
                F("data_assinatura", "Data da assinatura", "date", "01/12/2026", None, "Assinatura"),
            ],
        ),
    ]


GOLD: list[GoldDoc] = _docs()


def validate() -> list[str]:
    """Problemas de consistência do conjunto (vazio = ok)."""
    from app.modules.fields.validators import is_valid_cnpj, is_valid_cpf, normalize_oab

    problems: list[str] = []
    if len(GOLD) != 15:
        problems.append(f"{len(GOLD)} documentos; o plano pede 15")
    if missing := set(CATEGORIES) - {d.category for d in GOLD}:
        problems.append(f"categorias sem documento: {sorted(missing)}")
    for d in GOLD:
        names = [f.name for f in d.fields]
        if len(set(names)) != len(names):
            problems.append(f"{d.id}: campo repetido")
        used = {m for p in [d.title, *d.paragraphs, d.header or ""] for m in _MARK.findall(p)}
        if used ^ set(names):
            problems.append(f"{d.id}: marcadores e campos não batem: {sorted(used ^ set(names))}")
        for f in d.fields:
            if f.type not in FIELD_TYPES:
                problems.append(f"{d.id}.{f.name}: tipo {f.type}")
            if f.type == "cpf" and not is_valid_cpf(f.value):
                problems.append(f"{d.id}.{f.name}: CPF inválido")
            if f.type == "cnpj" and not is_valid_cnpj(f.value):
                problems.append(f"{d.id}.{f.name}: CNPJ inválido")
            if f.type == "oab" and not normalize_oab(f.value):
                problems.append(f"{d.id}.{f.name}: OAB inválida")
        blocks = dict(d.render())
        for o in d.occurrences:
            if blocks[o.block][o.start : o.end] != o.text:
                problems.append(f"{d.id}: ocorrência fora do lugar em {o.block}")
    return problems
