"""DOCX de exemplo (produto: "experimentar com um exemplo"; também seed, E2E e testes). Nenhum dado pessoal
real: nomes, CPF e CNPJ são fictícios/válidos apenas em dígitos verificadores.

`build_demo_contract()` — "Contrato de prestação de serviços":
  - texto com formatação (negrito, itálico), título centralizado;
  - "João da Silva" dividido em vários runs ("Jo" + "ão da Silva") e repetido em 4 lugares;
  - tabela, cabeçalho, rodapé, imagem inline e quebra de página.
`build_demo_receipt()` — "Recibo de pagamento": mesma pessoa, CPF e valor do contrato, para o kit de exemplo
mostrar variáveis compartilhadas entre documentos diferentes.
"""

from __future__ import annotations

import io
import struct
import zlib

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Cm, Pt


def tiny_png(width: int = 96, height: int = 32, rgb: tuple[int, int, int] = (30, 64, 175)) -> bytes:
    """PNG sólido gerado sem Pillow (uso em logo de demonstração)."""

    def chunk(kind: bytes, data: bytes) -> bytes:
        return (
            struct.pack(">I", len(data))
            + kind
            + data
            + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)
        )

    raw = b"".join(b"\x00" + bytes(rgb) * width for _ in range(height))
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw, 9))
        + chunk(b"IEND", b"")
    )


def build_demo_contract() -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.5)

    header_p = section.header.paragraphs[0]
    header_p.text = "Estúdio Exemplo Serviços Ltda. — Contrato de Prestação de Serviços"
    header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_p = section.footer.paragraphs[0]
    footer_p.text = "Documento de demonstração — contratante: João da Silva"
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("CONTRATO DE PRESTAÇÃO DE SERVIÇOS")
    r.bold = True
    r.font.size = Pt(16)

    # imagem inline (logo) em parágrafo próprio
    logo_p = doc.add_paragraph()
    logo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo_p.add_run().add_picture(io.BytesIO(tiny_png()), width=Cm(3))

    p = doc.add_paragraph()
    r = p.add_run("CONTRATANTE: ")
    r.bold = True
    p.add_run("Jo")  # o nome fica dividido em runs, como o Word costuma fazer
    p.add_run("ão da Silva")
    p.add_run(", inscrito no CPF sob o nº ")
    p.add_run("123.456.789-09")
    p.add_run(", doravante denominado CONTRATANTE.")

    p = doc.add_paragraph()
    r = p.add_run("CONTRATADA: ")
    r.bold = True
    p.add_run("Estúdio Exemplo Serviços Ltda., inscrita no CNPJ sob o nº ")
    p.add_run("11.222.333/0001-81")
    p.add_run(", doravante denominada CONTRATADA.")

    p = doc.add_paragraph()
    r = p.add_run("CLÁUSULA 1ª – DO OBJETO. ")
    r.bold = True
    p.add_run("O presente contrato tem por objeto a prestação de serviços de ")
    r = p.add_run("consultoria em desenvolvimento de software")
    r.italic = True
    p.add_run(", a ser executada pela CONTRATADA em favor de João da Silva.")

    p = doc.add_paragraph()
    r = p.add_run("CLÁUSULA 2ª – DO VALOR. ")
    r.bold = True
    p.add_run("Pela prestação dos serviços, o CONTRATANTE pagará à CONTRATADA o valor mensal de ")
    p.add_run("R$ 5.000,00")
    p.add_run(".")

    p = doc.add_paragraph()
    r = p.add_run("CLÁUSULA 3ª – DA VIGÊNCIA. ")
    r.bold = True
    p.add_run("O contrato inicia-se em ")
    p.add_run("01/10/2026")
    p.add_run(" e terá vigência de 12 meses. Renovação automática: ")
    p.add_run("Sim")
    p.add_run(".")

    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    rows = [("Item", "Detalhe"), ("Contratante", "João da Silva"), ("Valor mensal", "R$ 5.000,00")]
    for i, (a, b) in enumerate(rows):
        table.cell(i, 0).text = a
        table.cell(i, 1).text = b
        if i == 0:
            for c in table.rows[0].cells:
                for run in c.paragraphs[0].runs:
                    run.bold = True

    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)

    p = doc.add_paragraph("São Paulo, 21 de setembro de 2026.")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_paragraph("_______________________________")
    p = doc.add_paragraph()
    p.add_run("João da Silva").bold = True
    p.add_run(" — CONTRATANTE")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_demo_receipt() -> bytes:
    """Parágrafos do corpo, na ordem: p1 título, p2 texto (nome, CPF, valor, mês), p3 local e data,
    p4 linha de assinatura, p5 quem assina — `app.modules.samples.service` aplica os campos por essas chaves."""
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("RECIBO DE PAGAMENTO")
    r.bold = True
    r.font.size = Pt(16)

    p = doc.add_paragraph()
    p.add_run("Recebi de ")
    p.add_run("João da Silva").bold = True
    p.add_run(", inscrito no CPF sob o nº ")
    p.add_run("123.456.789-09")
    p.add_run(", a importância de ")
    p.add_run("R$ 5.000,00")
    p.add_run(", referente à mensalidade de ")
    p.add_run("outubro de 2026")
    p.add_run(" dos serviços prestados pela Estúdio Exemplo Serviços Ltda., dando plena quitação.")

    p = doc.add_paragraph("São Paulo, 05/10/2026.")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_paragraph("_______________________________")
    p = doc.add_paragraph()
    p.add_run("Estúdio Exemplo Serviços Ltda.").bold = True

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


DEMO_SERVICES = [
    ("Consultoria em desenvolvimento de software", "consultoria_software"),
    ("Suporte técnico mensal", "suporte_tecnico"),
    ("Treinamento de equipes", "treinamento"),
    ("Auditoria de processos", "auditoria"),
]
