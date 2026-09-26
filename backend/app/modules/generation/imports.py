"""Importação de .xlsx/.csv, sugestão de mapeamento de colunas e planilha-modelo com validações."""

from __future__ import annotations

import csv
import difflib
import io
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.core.errors import ValidationFailed
from app.docx_engine.safety import DocxRejected, validate_zip_container
from app.modules.fields.types import TYPE_LABELS, FieldDef, normalize_key

MAX_COLUMNS = 100


def _cell_to_str(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Sim" if v else "Não"
    if isinstance(v, datetime):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, date):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        return (
            str(int(v)) if v == int(v) else str(v).replace(".", ",")
        )  # pt-BR: o servidor interpreta vírgula como decimal
    return str(v).strip()


def _finish(matrix: list[list[str]], max_rows: int) -> tuple[list[str], list[list[str]]]:
    matrix = [r for r in matrix if any(c.strip() for c in r)]
    if not matrix:
        raise ValidationFailed("A planilha está vazia.", code="EMPTY_IMPORT")
    header, rows = matrix[0], matrix[1:]
    if len(header) > MAX_COLUMNS:
        raise ValidationFailed(f"A planilha tem colunas demais (máx. {MAX_COLUMNS}).", code="IMPORT_TOO_WIDE")
    if len(rows) > max_rows:
        raise ValidationFailed(
            f"A planilha tem {len(rows)} linhas; o limite por lote é {max_rows}.", code="IMPORT_TOO_LONG"
        )
    width = len(header)
    rows = [(r + [""] * width)[:width] for r in rows]
    header = [h.strip() or f"Coluna {i + 1}" for i, h in enumerate(header)]
    return header, rows


def parse_table(
    filename: str, data: bytes, *, max_rows: int, max_bytes: int = 10 * 1024 * 1024
) -> tuple[list[str], list[list[str]]]:
    name = filename.lower()
    if len(data) > max_bytes:
        raise ValidationFailed("O arquivo é grande demais (máx. 10 MB).", code="FILE_TOO_LARGE")
    if name.endswith(".xlsx"):
        try:
            validate_zip_container(data, max_bytes=max_bytes)
        except DocxRejected as exc:
            raise ValidationFailed(exc.message, code=exc.code) from exc
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            ws = wb["Modelo"] if "Modelo" in wb.sheetnames else wb.worksheets[0]
            matrix = [[_cell_to_str(c) for c in row] for row in ws.iter_rows(values_only=True)]
            wb.close()
        except Exception as exc:
            raise ValidationFailed(
                "Não foi possível ler a planilha. Verifique se é um .xlsx válido.", code="INVALID_XLSX"
            ) from exc
        return _finish(matrix, max_rows)
    if name.endswith(".csv") or name.endswith(".txt"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("cp1252", errors="replace")
        first = text.splitlines()[0] if text.strip() else ""
        delim = max((";", ",", "\t"), key=first.count)
        matrix = [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), delimiter=delim)]
        return _finish(matrix, max_rows)
    raise ValidationFailed(
        "Formato não suportado. Envie um arquivo .xlsx ou .csv.", code="INVALID_IMPORT_FORMAT"
    )


def suggest_mapping(columns: list[str], defs: list[FieldDef]) -> list[str | None]:
    """Para cada coluna do arquivo, o nome do campo mais provável (ou None)."""
    by_norm: dict[str, str] = {}
    for d in defs:
        by_norm.setdefault(normalize_key(d.label), d.name)
        by_norm.setdefault(normalize_key(d.name), d.name)
    used: set[str] = set()
    result: list[str | None] = []
    for col in columns:
        n = normalize_key(col)
        pick = by_norm.get(n)
        if pick is None and n:
            close = difflib.get_close_matches(n, list(by_norm), n=1, cutoff=0.82)
            pick = by_norm[close[0]] if close else None
        if pick in used:
            pick = None
        if pick:
            used.add(pick)
        result.append(pick)
    return result


def build_template_xlsx(generator_name: str, defs: list[FieldDef]) -> bytes:
    cols = [d for d in defs if d.include_in_batch]
    wb = Workbook()
    ws = wb.active
    ws.title = "Modelo"
    head_fill = PatternFill("solid", fgColor="1E3A8A")
    for i, d in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i, value=d.label + (" *" if d.required else ""))
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = max(16, min(40, len(d.label) + 6))
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B2"

    lists = wb.create_sheet("Listas")
    list_col = 0
    for i, d in enumerate(cols, start=1):
        rng = f"{get_column_letter(i)}2:{get_column_letter(i)}1001"
        dv: DataValidation | None = None
        if d.type == "select" and d.options:
            list_col += 1
            for r, o in enumerate(d.options, start=1):
                lists.cell(row=r, column=list_col, value=o.label)
            ref = f"Listas!${get_column_letter(list_col)}$1:${get_column_letter(list_col)}${len(d.options)}"
            dv = DataValidation(
                type="list",
                formula1=f"={ref}",
                allow_blank=not d.required,
                showErrorMessage=True,
                errorTitle="Valor inválido",
                error="Escolha uma das opções da lista.",
            )
        elif d.type in ("yes_no", "checkbox"):
            dv = DataValidation(
                type="list",
                formula1='"Sim,Não"',
                allow_blank=True,
                showErrorMessage=True,
                errorTitle="Valor inválido",
                error="Use Sim ou Não.",
            )
        elif d.type in ("number", "currency"):
            dv = DataValidation(
                type="decimal",
                operator="between",
                formula1="-999999999",
                formula2="999999999",
                allow_blank=True,
                showErrorMessage=True,
                errorTitle="Número inválido",
                error="Informe um número (ex.: 1234,56).",
            )
        elif d.type == "date":
            dv = DataValidation(
                type="date",
                operator="greaterThan",
                formula1="1",
                allow_blank=True,
                showErrorMessage=True,
                errorTitle="Data inválida",
                error="Informe uma data (dd/mm/aaaa).",
            )
            for r in range(2, 1002):
                ws.cell(row=r, column=i).number_format = "DD/MM/YYYY"
        elif d.type in ("cpf", "cnpj", "cep", "phone"):
            for r in range(2, 1002):
                ws.cell(row=r, column=i).number_format = "@"  # texto: preserva zeros à esquerda
        if dv is not None:
            ws.add_data_validation(dv)
            dv.add(rng)
    lists.sheet_state = "hidden"

    info = wb.create_sheet("Instruções")
    info.append([f"Planilha-modelo — {generator_name}"])
    info["A1"].font = Font(bold=True, size=14)
    info.append(["Preencha a aba “Modelo”: uma linha por documento. Colunas com * são obrigatórias."])
    info.append([])
    info.append(["Coluna", "Tipo", "Obrigatória", "Formato / exemplo", "Descrição"])
    for c in info[4]:
        c.font = Font(bold=True)
    examples = {
        "cpf": "123.456.789-09",
        "cnpj": "11.222.333/0001-81",
        "cep": "01310-100",
        "phone": "(11) 91234-5678",
        "date": "21/09/2026",
        "currency": "1.234,56",
        "number": "1.234,5",
        "email": "nome@empresa.com.br",
        "yes_no": "Sim ou Não",
        "checkbox": "Sim ou Não",
    }
    for d in cols:
        ex = examples.get(d.type, d.example or "")
        if d.type == "select":
            ex = "; ".join(o.label for o in d.options[:8]) + (" …" if len(d.options) > 8 else "")
        info.append(
            [
                d.label,
                TYPE_LABELS.get(d.type, d.type),
                "Sim" if d.required else "Não",
                ex,
                d.description or "",
            ]
        )
    for col, width in zip("ABCDE", (28, 18, 12, 44, 50), strict=True):
        info.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
