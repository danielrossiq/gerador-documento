"""Testes do motor DOCX: runs divididos, campos repetidos, formatação, imagem, tabela, cabeçalho/rodapé."""

from __future__ import annotations

import io
import re
import zipfile

import pytest
from docx import Document

from app.docx_engine.compile import BlockSpec
from app.docx_engine.inline import normalize_nodes
from app.docx_engine.nodes import replace_range_with_field, unwrap_field
from app.docx_engine.parse import parse_docx
from tests.helpers import all_text, compile_with, flat_text, offset_of, part_xml, render

F_NOME, F_CPF, F_VALOR = "fld-nome", "fld-cpf", "fld-valor"
NAMES = {F_NOME: "nome_contratante", F_CPF: "cpf_contratante", F_VALOR: "valor_mensal"}


def _name_fields(parsed):
    """Campo nome nas 5 ocorrências do 'João da Silva' (corpo, tabela, rodapé, assinatura)."""
    fields = {}
    for key in ("body/p3", "body/p5", "body/t1/r2/c2/p1", "body/p11", "footer1/p1"):
        s, e = offset_of(parsed, key, "João da Silva")
        fields[key] = [(s, e, F_NOME)]
    return fields


def test_original_has_name_split_across_runs(demo_docx):
    xml = part_xml(demo_docx)
    assert ">Jo<" in xml and ">ão da Silva<" in xml, "o fixture deve dividir o nome em dois runs"


def test_field_across_runs_is_replaced_without_naive_string_replace(demo_docx):
    parsed = parse_docx(demo_docx)
    s, e = offset_of(parsed, "body/p3", "João da Silva")
    _, res = compile_with(demo_docx, {"body/p3": [(s, e, F_NOME)]}, names=NAMES)
    assert res.tokens == {"nome_contratante": 1}
    out = render(res.data, {"nome_contratante": "Maria Souza"})
    p3 = next(t for t in all_text(out)["body"] if t.startswith("CONTRATANTE"))
    assert p3.startswith("CONTRATANTE: Maria Souza, inscrito no CPF")
    assert "João" not in p3


def test_repeated_field_replaced_everywhere(demo_docx):
    parsed = parse_docx(demo_docx)
    _, res = compile_with(demo_docx, _name_fields(parsed), names=NAMES)
    assert res.tokens["nome_contratante"] == 5
    out = render(res.data, {"nome_contratante": "Ana Lima"})
    text = all_text(out)
    assert "João da Silva" not in flat_text(out)
    assert text["tables"].count("Ana Lima") == 1
    assert "Ana Lima" in text["footer"][0]
    assert any("em favor de Ana Lima." in p for p in text["body"])
    assert any(p.startswith("Ana Lima") and p.endswith("CONTRATANTE") for p in text["body"])


def test_untouched_paragraphs_are_byte_identical(demo_docx):
    parsed = parse_docx(demo_docx)
    s, e = offset_of(parsed, "body/p3", "João da Silva")
    _, res = compile_with(demo_docx, {"body/p3": [(s, e, F_NOME)]}, names=NAMES)

    def paragraphs(data: bytes) -> list[str]:
        xml = part_xml(data)
        return re.findall(r"<w:p[ >].*?</w:p>", xml, flags=re.S)

    before, after = paragraphs(demo_docx), paragraphs(res.data)
    assert len(before) == len(after)
    diffs = [i for i, (a, b) in enumerate(zip(before, after, strict=True)) if a != b]
    assert diffs == [2], f"apenas o parágrafo do contratante (índice 2) deveria mudar; mudaram: {diffs}"


def test_formatting_is_inherited_by_field_run(demo_docx):
    parsed = parse_docx(demo_docx)
    s, e = offset_of(parsed, "body/p3", "CONTRATANTE", 0)  # "CONTRATANTE: " é run em negrito
    _, res = compile_with(demo_docx, {"body/p3": [(s, e, F_NOME)]}, names=NAMES)
    xml = part_xml(res.data)
    m = re.search(r"<w:r>(?:(?!</w:r>).)*?⟦F:deadbeef:nome_contratante⟧.*?</w:r>", xml, flags=re.S)
    assert m and "<w:b/>" in m.group(0), "o run do campo deve herdar o negrito do run de origem"


def test_bold_and_italic_toggle_by_marks(demo_docx):
    parsed = parse_docx(demo_docx)
    block = next(b for b in parsed.blocks if b.key == "body/p7")
    content = [dict(n) for n in block.content]
    # coloca em negrito o trecho "12 meses" (originalmente sem negrito) via marcas
    text = block.text
    s = text.index("12 meses")
    parts = []
    pos = 0
    for n in content:
        t = n["text"]
        if pos <= s < pos + len(t):
            a = s - pos
            parts += [
                {"type": "text", "text": t[:a], **({"marks": n["marks"]} if n.get("marks") else {})},
                {"type": "text", "text": "12 meses", "marks": [{"type": "bold"}]},
                {"type": "text", "text": t[a + 8 :], **({"marks": n["marks"]} if n.get("marks") else {})},
            ]
        else:
            parts.append(n)
        pos += len(t)
    _, res = compile_with(demo_docx, overrides={"body/p7": normalize_nodes(parts)})
    doc = Document(io.BytesIO(res.data))
    p = next(p for p in doc.paragraphs if p.text.startswith("CLÁUSULA 3ª"))
    bold_runs = [r.text for r in p.runs if r.bold]
    assert any("12 meses" in t for t in bold_runs)
    assert "CLÁUSULA 3ª – DA VIGÊNCIA. " in bold_runs  # negrito original preservado
    assert p.text == next(b.text for b in parsed.blocks if b.key == "body/p7")


def test_image_is_preserved_through_compile_and_render(demo_docx):
    parsed = parse_docx(demo_docx)
    _, res = compile_with(demo_docx, _name_fields(parsed), names=NAMES)
    out = render(res.data, {"nome_contratante": "Ana Lima"})
    z = zipfile.ZipFile(io.BytesIO(out))
    assert any(n.startswith("word/media/") for n in z.namelist())
    assert "<w:drawing>" in part_xml(out)
    assert len(Document(io.BytesIO(out)).inline_shapes) == 1


def test_rendered_docx_is_a_valid_document(demo_docx):
    parsed = parse_docx(demo_docx)
    _, res = compile_with(demo_docx, _name_fields(parsed), names=NAMES)
    out = render(res.data, {"nome_contratante": "Ana Lima"})
    doc = Document(io.BytesIO(out))  # abre sem erro
    assert len(doc.paragraphs) >= 10
    assert zipfile.ZipFile(io.BytesIO(out)).testzip() is None


def test_edit_plain_text_keeps_formatting(demo_docx):
    parsed = parse_docx(demo_docx)
    block = next(b for b in parsed.blocks if b.key == "body/p7")
    new = []
    for n in block.content:
        n = dict(n)
        n["text"] = n["text"].replace("12 meses", "24 meses")
        new.append(n)
    _, res = compile_with(demo_docx, overrides={"body/p7": new})
    doc = Document(io.BytesIO(res.data))
    p = next(p for p in doc.paragraphs if p.text.startswith("CLÁUSULA 3ª"))
    assert "24 meses" in p.text and "12 meses" not in p.text
    assert p.runs[0].bold is True


def test_multiline_and_special_characters_in_values(demo_docx):
    parsed = parse_docx(demo_docx)
    s, e = offset_of(parsed, "body/p3", "João da Silva")
    _, res = compile_with(demo_docx, {"body/p3": [(s, e, F_NOME)]}, names=NAMES)
    out = render(res.data, {"nome_contratante": 'Linha 1\nLinha 2 & <tag> "aspas"'})
    xml = part_xml(out)
    assert "<w:br/>" in xml
    assert "&amp;" in xml and "&lt;tag&gt;" in xml
    assert 'Linha 2 & <tag> "aspas"' in flat_text(out)


def test_insert_and_delete_paragraphs(demo_docx):
    from app.docx_engine.compile import compile_docx
    from tests.helpers import specs_from_parsed

    parsed = parse_docx(demo_docx)
    # remove body/p10 (linha de assinatura) e insere um parágrafo novo logo após body/p3 (ordinal 2)
    specs = [s for s in specs_from_parsed(parsed) if s.key != "body/p10"]
    for s in specs:
        if s.container_key == "body" and s.ordinal >= 3:
            s.ordinal += 1
    extra = BlockSpec(
        "body/n1",
        "body",
        3,
        "paragraph",
        "inserted",
        [{"type": "text", "text": "Parágrafo novo inserido."}],
        {"clone_of": "body/p4"},
    )
    res = compile_docx(demo_docx, [*specs, extra], {}, nonce="deadbeef")
    assert (res.inserted_blocks, res.deleted_blocks) == (1, 1)
    body = all_text(res.data)["body"]
    assert "Parágrafo novo inserido." in body
    assert "_______________________________" not in body
    idx_contratante = body.index(next(t for t in body if t.startswith("CONTRATANTE:")))
    assert body.index("Parágrafo novo inserido.") == idx_contratante + 1


def test_unwrap_field_keeps_or_drops_text():
    content = [
        {"type": "text", "text": "Olá "},
        {"type": "field", "attrs": {"fieldId": "x", "sample": "Maria"}},
        {"type": "text", "text": "!"},
    ]
    kept, n = unwrap_field(content, "x", keep_text=True)
    assert n == 1 and kept == [{"type": "text", "text": "Olá Maria!"}]
    dropped, _ = unwrap_field(content, "x", keep_text=False)
    assert dropped == [{"type": "text", "text": "Olá !"}]


def test_replace_range_rejects_non_text_and_out_of_bounds():
    content = [{"type": "text", "text": "abc"}, {"type": "hardBreak"}, {"type": "text", "text": "def"}]
    with pytest.raises(ValueError):
        replace_range_with_field(content, 2, 5, "f")
    with pytest.raises(ValueError):
        replace_range_with_field(content, 0, 99, "f")
    new, sample = replace_range_with_field(content, 0, 2, "f")
    assert sample == "ab" and new[0]["type"] == "field"


def test_insert_field_at_cursor_without_selection():
    content = [{"type": "text", "text": "Nome: ", "marks": [{"type": "bold"}]}]
    new, sample = replace_range_with_field(content, 6, 6, "f")
    assert sample == "" and new[-1]["type"] == "field"
    assert new[-1]["marks"] == [{"type": "bold"}]


def test_field_can_cross_runs_with_different_formatting(demo_docx):
    """'CONTRATANTE: João' cobre um run em negrito e um normal: o campo vira um run só."""
    parsed = parse_docx(demo_docx)
    s, e = offset_of(parsed, "body/p3", "CONTRATANTE: João")
    _, res = compile_with(demo_docx, {"body/p3": [(s, e, F_NOME)]}, names=NAMES)
    out = render(res.data, {"nome_contratante": "XYZ"})
    assert next(t for t in all_text(out)["body"] if t.startswith("XYZ")).startswith("XYZ da Silva, inscrito")


def test_document_without_tokens_renders_identically(demo_docx):
    _, res = compile_with(demo_docx)
    assert res.tokens == {} and res.modified_blocks == 0
    assert flat_text(render(res.data, {})) == flat_text(demo_docx)
