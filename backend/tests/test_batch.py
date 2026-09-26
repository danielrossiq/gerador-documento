"""Lote: validação, falha parcial, idempotência, ZIP + relatório, importação XLSX/CSV e planilha-modelo."""

from __future__ import annotations

import csv
import io
import zipfile

from openpyxl import Workbook, load_workbook

from tests.flows import SAMPLE_VALUES, build_generator, docx_text


def _row(i: int, **over) -> dict:
    base = dict(SAMPLE_VALUES, nome_do_contratante=f"Cliente {i:02d}")
    base.update(over)
    return base


def _ready(api, demo_docx):
    api.register()
    return build_generator(api, demo_docx)["id"]


def _zip_of(api, job) -> zipfile.ZipFile:
    z = next(f for f in job["files"] if f["kind"] == "zip")
    url = api.post(f"/files/{z['id']}/link").json()["url"]
    return zipfile.ZipFile(io.BytesIO(api.tc.get(url).content))


def test_batch_generates_zip_and_report(api, demo_docx):
    gid = _ready(api, demo_docx)
    rows = [_row(i) for i in range(1, 4)]
    r = api.post(
        f"/generators/{gid}/batches", json={"rows": rows, "format": "docx", "idempotency_key": "b-1"}
    )
    assert r.status_code == 202, r.text
    job = r.json()
    assert (job["status"], job["total_items"], job["succeeded_items"], job["failed_items"]) == (
        "succeeded",
        3,
        3,
        0,
    )
    z = _zip_of(api, job)
    names = sorted(z.namelist())
    assert len(names) == 4 and "relatorio.csv" in names  # 3 documentos + relatório
    docs = [n for n in names if n.endswith(".docx")]
    assert len(docs) == 3 and len({d.lower() for d in docs}) == 3
    texts = [docx_text(z.read(d)) for d in docs]
    for i in range(1, 4):
        assert any(f"Cliente {i:02d}" in t for t in texts)
    report = z.read("relatorio.csv").decode("utf-8-sig")
    rows_csv = list(csv.reader(io.StringIO(report), delimiter=";"))
    assert rows_csv[0] == ["linha", "status", "arquivos", "codigo_erro", "mensagem"]
    assert [r[1] for r in rows_csv[1:]] == ["gerado"] * 3
    # relatório também fica disponível separadamente
    assert any(f["kind"] == "report" for f in job["files"])


def test_invalid_rows_block_generation_or_can_be_skipped_explicitly(api, demo_docx):
    gid = _ready(api, demo_docx)
    rows = [_row(1), _row(2, cpf_do_contratante="000.000.000-00"), _row(3)]
    val = api.post(f"/generators/{gid}/batch/validate", json={"rows": rows}).json()
    assert (val["valid_count"], val["invalid_count"]) == (2, 1)
    assert val["rows"][1]["errors"] == {"cpf_do_contratante": "CPF inválido."}

    blocked = api.post(f"/generators/{gid}/batches", json={"rows": rows})
    assert blocked.status_code == 422 and blocked.json()["error"]["code"] == "BATCH_INVALID_ROWS"
    assert blocked.json()["error"]["details"]["invalid_count"] == 1
    assert api.get("/jobs").json()["total"] == 0  # nada foi criado

    ok = api.post(f"/generators/{gid}/batches", json={"rows": rows, "only_valid": True}).json()
    assert ok["status"] == "partial" and ok["succeeded_items"] == 2 and ok["skipped_items"] == 1
    items = api.get(f"/jobs/{ok['id']}/items").json()["items"]
    skipped = next(i for i in items if i["status"] == "skipped")
    assert skipped["row_index"] == 2 and "CPF" in skipped["error_message"]
    report = _zip_of(api, ok).read("relatorio.csv").decode("utf-8-sig")
    assert "ignorada" in report and "CPF inválido" in report


def test_partial_failure_does_not_stop_the_batch(api, demo_docx, fake_pdf):
    """Uma linha falha na conversão para PDF; as outras concluem e o job fica 'partial' (falha isolada)."""
    fake_pdf.fail_for = {"Cliente 02"}
    gid = _ready(api, demo_docx)
    rows = [_row(i) for i in range(1, 5)]
    job = api.post(f"/generators/{gid}/batches", json={"rows": rows, "format": "pdf"}).json()
    assert job["status"] == "partial"
    assert (job["succeeded_items"], job["failed_items"]) == (3, 1)
    items = {i["row_index"]: i for i in api.get(f"/jobs/{job['id']}/items").json()["items"]}
    assert items[2]["status"] == "failed" and items[2]["error_code"] == "PDF_CONVERSION_FAILED"
    assert all(items[i]["status"] == "succeeded" and items[i]["files"][0]["kind"] == "pdf" for i in (1, 3, 4))
    report = _zip_of(api, job).read("relatorio.csv").decode("utf-8-sig")
    assert "PDF_CONVERSION_FAILED" in report and report.count("gerado") == 3


def test_all_items_failing_marks_job_failed(api, demo_docx, fake_pdf):
    fake_pdf.fail_for = {"Cliente"}
    gid = _ready(api, demo_docx)
    job = api.post(f"/generators/{gid}/batches", json={"rows": [_row(1), _row(2)], "format": "pdf"}).json()
    assert (
        job["status"] == "failed" and job["error_code"] == "ALL_ITEMS_FAILED" and job["succeeded_items"] == 0
    )


def test_idempotency_key_prevents_duplicate_batches(api, demo_docx):
    gid = _ready(api, demo_docx)
    body = {"rows": [_row(1), _row(2)], "format": "docx"}
    first = api.post(f"/generators/{gid}/batches", json=body, headers={"Idempotency-Key": "abc-123"})
    again = api.post(f"/generators/{gid}/batches", json=body, headers={"Idempotency-Key": "abc-123"})
    assert first.status_code == 202 and again.status_code == 200
    assert again.headers["idempotent-replayed"] == "true" and first.headers["idempotent-replayed"] == "false"
    assert first.json()["id"] == again.json()["id"]
    assert api.get("/jobs").json()["total"] == 1
    other = api.post(f"/generators/{gid}/batches", json=body, headers={"Idempotency-Key": "abc-999"})
    assert other.json()["id"] != first.json()["id"]
    # também vale para a geração individual
    s1 = api.post(
        f"/generators/{gid}/generate", json={"values": SAMPLE_VALUES}, headers={"Idempotency-Key": "s-1"}
    ).json()
    s2 = api.post(
        f"/generators/{gid}/generate", json={"values": SAMPLE_VALUES}, headers={"Idempotency-Key": "s-1"}
    ).json()
    assert s1["id"] == s2["id"]


def test_rerunning_a_finished_job_is_a_noop_and_resumes_when_interrupted(api, demo_docx):
    import uuid
    from datetime import timedelta

    from sqlalchemy import func, select, update

    from app.db.base import utcnow
    from app.db.models import GeneratedFile, GenerationItem, GenerationJob
    from app.db.session import get_factory
    from app.modules.generation.runner import run_generation_job

    gid = _ready(api, demo_docx)
    job = api.post(f"/generators/{gid}/batches", json={"rows": [_row(i) for i in range(1, 4)]}).json()
    jid = uuid.UUID(job["id"])

    def n_files() -> int:
        with get_factory()() as db:
            return db.scalar(select(func.count(GeneratedFile.id)).where(GeneratedFile.job_id == jid))

    before = n_files()
    run_generation_job(job["id"])  # job já terminal: não faz nada
    assert n_files() == before

    # simula queda no meio: job ficou 'processing' com heartbeat velho (executor morto) e um item pendente
    # → retoma sem duplicar
    with get_factory()() as db:
        db.execute(
            update(GenerationJob)
            .where(GenerationJob.id == jid)
            .values(status="processing", finished_at=None, heartbeat_at=utcnow() - timedelta(hours=1))
        )
        item = db.scalar(
            select(GenerationItem).where(GenerationItem.job_id == jid, GenerationItem.row_index == 3)
        )
        item.status = "pending"
        db.commit()
    run_generation_job(job["id"])
    done = api.get(f"/jobs/{job['id']}").json()
    assert done["status"] == "succeeded" and done["succeeded_items"] == 3
    with get_factory()() as db:
        docx_per_item = db.execute(
            select(GeneratedFile.item_id, func.count())
            .where(GeneratedFile.kind == "docx")
            .group_by(GeneratedFile.item_id)
        ).all()
    assert all(n == 1 for _, n in docx_per_item)  # nenhum documento duplicado


def test_transient_storage_failure_keeps_job_retryable(api, demo_docx, monkeypatch):
    """Falha transitória no storage ao gravar um documento gerado (não ao preparar o modelo): com backend sem
    fila de retry, o estado é terminal e explícito."""
    from app.adapters.storage import StorageError, get_storage

    gid = _ready(api, demo_docx)
    # 1ª geração bem-sucedida: compila e cacheia o gabarito, então a próxima chamada não escreve no storage
    # antes de criar o job (isola a falha no passo que este teste quer cobrir: gravar o documento da linha).
    assert api.post(f"/generators/{gid}/batches", json={"rows": [_row(0)]}).json()["status"] == "succeeded"

    storage = get_storage()
    real_put = storage.put
    monkeypatch.setattr(storage, "put", lambda *a, **k: (_ for _ in ()).throw(StorageError("indisponível")))
    monkeypatch.setattr("app.modules.generation.runner.time.sleep", lambda s: None)
    job = api.post(f"/generators/{gid}/batches", json={"rows": [_row(1)]}).json()
    assert job["status"] == "failed" and job["error_code"] == "TRANSIENT_ERROR"
    monkeypatch.setattr(storage, "put", real_put)


def test_transient_storage_failure_while_preparing_template_is_a_clear_503(api, demo_docx, monkeypatch):
    """A 1ª compilação do gabarito também grava no storage; se falhar, nenhum job chega a ser criado."""
    from app.adapters.storage import StorageError, get_storage

    gid = _ready(api, demo_docx)
    storage = get_storage()
    monkeypatch.setattr(storage, "put", lambda *a, **k: (_ for _ in ()).throw(StorageError("indisponível")))
    monkeypatch.setattr("app.modules.templates.service.time.sleep", lambda s: None)
    r = api.post(f"/generators/{gid}/batches", json={"rows": [_row(1)]})
    assert r.status_code == 503 and r.json()["error"]["code"] == "TEMPLATE_UNAVAILABLE"
    assert api.get("/jobs").json()["total"] == 0


def test_filenames_are_sanitized_and_deduplicated(api, demo_docx):
    gid = _ready(api, demo_docx)
    api.patch(
        f"/generators/{gid}", json={"filename_pattern": "Contrato_{nome_do_contratante}_{data_de_inicio}"}
    )
    rows = [
        _row(1, nome_do_contratante="Ana/Lima: <teste>"),
        _row(2, nome_do_contratante="ana/lima: <teste>"),
        _row(3, nome_do_contratante="CON"),
    ]
    job = api.post(f"/generators/{gid}/batches", json={"rows": rows}).json()
    names = sorted(i["filename"] for i in api.get(f"/jobs/{job['id']}/items").json()["items"])
    assert names[0].startswith("Contrato_Ana-Lima")
    assert len({n.lower() for n in names}) == 3 and all(
        "/" not in n and ":" not in n and "<" not in n for n in names
    )
    assert any(n.endswith("_2") for n in names)


def test_limits_are_enforced_server_side(api, demo_docx):
    gid = _ready(api, demo_docx)
    too_many = api.post(f"/generators/{gid}/batches", json={"rows": [_row(1)] * 1001})
    assert too_many.status_code == 429 and too_many.json()["error"]["code"] == "BATCH_TOO_LARGE"
    assert api.post(f"/generators/{gid}/batches", json={"rows": []}).json()["error"]["code"] == "EMPTY_BATCH"


def test_thousand_rows_validate_quickly(api, demo_docx):
    import time

    gid = _ready(api, demo_docx)
    rows = [_row(i) for i in range(1000)]
    t0 = time.monotonic()
    val = api.post(f"/generators/{gid}/batch/validate", json={"rows": rows}).json()
    assert val["valid_count"] == 1000 and time.monotonic() - t0 < 10


def test_thousand_documents_batch_completes(api, demo_docx):
    """Arquitetura suporta 1.000 documentos por lote (DOCX). Mede tempo; PDF é medido no container do worker."""
    import time

    gid = _ready(api, demo_docx)
    rows = [_row(i) for i in range(1, 1001)]
    t0 = time.monotonic()
    job = api.post(f"/generators/{gid}/batches", json={"rows": rows, "format": "docx"}).json()
    elapsed = time.monotonic() - t0
    assert job["status"] == "succeeded" and job["succeeded_items"] == 1000
    z = _zip_of(api, job)
    assert len([n for n in z.namelist() if n.endswith(".docx")]) == 1000
    print(f"\n1000 documentos DOCX gerados em {elapsed:.1f}s")


def test_repeat_and_cancel(api, demo_docx):
    gid = _ready(api, demo_docx)
    first = api.post(f"/generators/{gid}/batches", json={"rows": [_row(1), _row(2)]}).json()
    again = api.post(f"/jobs/{first['id']}/repeat").json()
    assert (
        again["id"] != first["id"] and again["source_job_id"] == first["id"] and again["succeeded_items"] == 2
    )
    assert api.post(f"/jobs/{again['id']}/cancel").status_code == 409  # já concluído


# ------------------------------------------------------------------ importação e planilha-modelo


def _xlsx(header: list[str], rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(header)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_import_xlsx_with_mapping_and_error_detection(api, demo_docx):
    import datetime

    gid = _ready(api, demo_docx)
    data = _xlsx(
        ["Nome do Contratante *", "CPF", "Serviço", "Valor (R$)", "Início", "Coluna extra"],
        [
            [
                "Ana Lima",
                "529.982.247-25",
                "Suporte técnico mensal",
                1500.5,
                datetime.date(2026, 10, 1),
                "ignorar",
            ],
            ["Bruno Dias", "111.111.111-11", "Treinamento de equipes", 800, datetime.date(2026, 11, 3), "x"],
        ],
    )
    up = api.post(
        f"/generators/{gid}/batch/imports", files={"file": ("lote.xlsx", data, "application/octet-stream")}
    )
    assert up.status_code == 201, up.text
    body = up.json()
    assert body["row_count"] == 2 and body["columns"][0] == "Nome do Contratante *"
    # mapeamento sugerido: reconhece "Nome do Contratante *" pelo rótulo; "Coluna extra" fica sem campo
    assert body["suggested_mapping"][0] == "nome_do_contratante" and body["suggested_mapping"][5] is None
    # cada campo vem com id (pra permitir PUT .../options) e, se for lista suspensa, as opções atuais
    by_name = {f["name"]: f for f in body["fields"]}
    assert by_name["nome_do_contratante"]["id"] and by_name["nome_do_contratante"]["options"] is None
    assert by_name["tipo_de_servico"]["options"] and all(
        {"label", "value"} <= set(o) for o in by_name["tipo_de_servico"]["options"]
    )
    mapping = [
        "nome_do_contratante",
        "cpf_do_contratante",
        "tipo_de_servico",
        "valor_mensal",
        "data_de_inicio",
        None,
    ]
    applied = api.post(
        f"/generators/{gid}/batch/imports/{body['import_id']}/apply", json={"mapping": mapping}
    ).json()
    assert (
        applied["rows"][0]["valor_mensal"] == "1500,5"
    )  # número do Excel → texto pt-BR (o servidor interpreta a vírgula)
    assert applied["rows"][0]["data_de_inicio"] == "01/10/2026"
    v = applied["validation"]
    assert (v["valid_count"], v["invalid_count"]) == (1, 1) or v["invalid_count"] >= 1
    assert "cpf_do_contratante" in v["rows"][1]["errors"]
    # o import é descartado após aplicar
    assert (
        api.post(
            f"/generators/{gid}/batch/imports/{body['import_id']}/apply", json={"mapping": mapping}
        ).status_code
        == 422
    )


def test_import_csv_semicolon_and_cp1252(api, demo_docx):
    gid = _ready(api, demo_docx)
    csv_bytes = "Nome do contratante;CPF do contratante\nJoão Ávila;529.982.247-25\n".encode("cp1252")
    up = api.post(
        f"/generators/{gid}/batch/imports", files={"file": ("lote.csv", csv_bytes, "text/csv")}
    ).json()
    assert (
        up["columns"] == ["Nome do contratante", "CPF do contratante"] and up["sample"][0][0] == "João Ávila"
    )
    assert up["suggested_mapping"] == ["nome_do_contratante", "cpf_do_contratante"]


def test_import_rejects_bad_files_and_duplicate_mapping(api, demo_docx):
    gid = _ready(api, demo_docx)
    bad = api.post(
        f"/generators/{gid}/batch/imports", files={"file": ("x.exe", b"MZ", "application/octet-stream")}
    )
    assert bad.status_code == 422 and bad.json()["error"]["code"] == "INVALID_IMPORT_FORMAT"
    fake = api.post(
        f"/generators/{gid}/batch/imports",
        files={"file": ("x.xlsx", b"nao e zip", "application/octet-stream")},
    )
    assert fake.status_code == 422
    up = api.post(
        f"/generators/{gid}/batch/imports", files={"file": ("a.csv", b"a;b\n1;2\n", "text/csv")}
    ).json()
    dup = api.post(
        f"/generators/{gid}/batch/imports/{up['import_id']}/apply",
        json={"mapping": ["cpf_do_contratante", "cpf_do_contratante"]},
    )
    assert dup.json()["error"]["code"] == "INVALID_MAPPING"


def test_template_xlsx_has_headers_validations_and_instructions(api, demo_docx):
    gid = _ready(api, demo_docx)
    r = api.get(f"/generators/{gid}/batch/template.xlsx")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Modelo"]
    headers = [c.value for c in ws[1]]
    assert (
        headers[0] == "Nome do contratante *"
        and "Tipo de serviço *" in headers
        and "CNPJ da contratada" in headers
    )
    assert len(ws.data_validations.dataValidation) >= 3  # listas, decimais, datas…
    info = wb["Instruções"]
    assert any(
        row[0] == "Tipo de serviço" and "Suporte técnico mensal" in (row[3] or "")
        for row in info.iter_rows(values_only=True)
    )
