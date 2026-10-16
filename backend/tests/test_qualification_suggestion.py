"""Qualificação das partes: estado civil e regime de bens encontrados no documento viram listas prontas
(no gênero e na caixa do texto original), por parte. Regras locais — sem rede e sem IA."""

from __future__ import annotations

import io
import socket

import pytest
from docx import Document

from app.modules.suggestions.rules import detect_rules
from tests.flows import docx_text, upload


def _selects(text: str):
    return [h for h in detect_rules([("b", text)]) if h.type == "select" and h.name != "mes"]


def _labels(h) -> list[str]:
    return [o["label"] for o in h.config["options"]]


@pytest.fixture
def no_network(monkeypatch: pytest.MonkeyPatch):
    def blocked(*_a, **_k):
        raise AssertionError("tentativa de conexão de rede")

    monkeypatch.setattr(socket.socket, "connect", blocked)
    monkeypatch.setattr(socket, "create_connection", blocked)
    monkeypatch.setattr(socket, "getaddrinfo", blocked)


def test_feminine_civil_status_and_property_regime_with_party_role(no_network):
    civil, regime = _selects(
        "CONTRATANTE: Maria Souza, brasileira, casada sob o regime de comunhão parcial de bens, empresária"
    )
    assert (civil.text, civil.name, civil.label) == (
        "casada",
        "estado_civil_contratante",
        "Estado civil — Contratante",
    )
    assert _labels(civil) == [
        "solteira",
        "casada",
        "divorciada",
        "viúva",
        "separada judicialmente",
        "em união estável",
    ]
    assert regime.text == "comunhão parcial de bens" and regime.name == "regime_de_bens_contratante"
    assert "comunhão parcial de bens" in _labels(regime) and "participação final nos aquestos" in _labels(
        regime
    )


def test_masculine_list_and_two_parties_get_different_fields(no_network):
    hits = _selects(
        "OUTORGANTE: João Lima, brasileiro, solteiro. OUTORGADO: Pedro Alves, brasileiro, divorciado."
    )
    assert [(h.text, h.name) for h in hits] == [
        ("solteiro", "estado_civil_outorgante"),
        ("divorciado", "estado_civil_outorgado"),
    ]
    assert _labels(hits[0])[:2] == ["solteiro", "casado"]


def test_capitalized_term_gets_capitalized_options(no_network):
    (h,) = _selects("Casado, residente em Serra/ES")
    assert _labels(h)[:2] == ["Solteiro", "Casado"]


@pytest.mark.parametrize(
    "text", ["foi separado do processo", "casamento marcado", "a solteirice", "o estado"]
)
def test_ordinary_words_are_not_qualification(text, no_network):
    assert _selects(text) == []


def test_upload_accept_and_generate_with_another_option(api, monkeypatch: pytest.MonkeyPatch):
    # sem provedor de IA: só as regras locais analisam o documento enviado
    monkeypatch.setattr("app.modules.suggestions.service.get_ai_provider", lambda: None)
    api.register()
    gid = api.post("/generators", json={"name": "Procuração"}).json()["id"]
    doc = Document()
    doc.add_paragraph("PROCURAÇÃO AD JUDICIA")
    doc.add_paragraph("OUTORGANTE: Maria Souza, brasileira, casada, empresária, CPF 529.982.247-25.")
    buf = io.BytesIO()
    doc.save(buf)
    assert upload(api, gid, buf.getvalue()).status_code in (200, 201)

    items = api.get(f"/generators/{gid}/suggestions").json()["items"]
    (civil,) = [s for s in items if s["suggested_name"] == "estado_civil_outorgante"]
    field = api.post(f"/suggestions/{civil['id']}/accept").json()["field"]
    assert field["type"] == "select" and len(field["options"]) == 6
    assert "options" not in (field["config"] or {})

    job = api.post(
        f"/generators/{gid}/generate", json={"values": {field["name"]: "divorciada"}, "format": "docx"}
    ).json()
    (file,) = job["files"]
    text = docx_text(api.tc.get(api.post(f"/files/{file['id']}/link").json()["url"]).content)
    assert "brasileira, divorciada, empresária" in text and "casada" not in text
