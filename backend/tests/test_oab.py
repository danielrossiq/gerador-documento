"""Tipo de campo OAB: validação/formatação (vetores compartilhados cobrem as variações), sugestão automática,
geração e a migration 0005 (upgrade/downgrade) em PostgreSQL real quando MIGRATION_DATABASE_URL estiver definida."""

from __future__ import annotations

import os

import pytest
import sqlalchemy as sa

from app.modules.suggestions.rules import detect_rules
from tests.flows import docx_text, upload


def _oab(text: str):
    return [(h.text, h.name) for h in detect_rules([("b", text)]) if h.type == "oab"]


def test_rule_finds_oab_with_or_without_dots_and_role():
    assert _oab("Dr. Pedro Alves, advogado, OAB/ES 12.345, com escritório") == [("OAB/ES 12.345", "oab")]
    assert _oab("OUTORGADO: Ana Lima, advogada, OAB/SP nº 123.456-A.") == [
        ("OAB/SP nº 123.456-A", "oab_outorgado")
    ]
    assert _oab("OAB/ES 12345 e OAB/RJ 1.234") == [("OAB/ES 12345", "oab"), ("OAB/RJ 1.234", "oab")]


@pytest.mark.parametrize(
    "text",
    [
        "a OAB/XX 123 não existe",  # UF inválida
        "OAB/ES 1234567",  # mais de 6 dígitos
        "inscrita na OAB/SP sob o nº 123.456",  # trocar o trecho mudaria a redação
        "a OAB publicou",
    ],
)
def test_rule_ignores_what_is_not_a_registration(text):
    assert _oab(text) == []


def test_oab_field_via_api_validates_and_generates(api, demo_docx):
    api.register()
    gid = api.post("/generators", json={"name": "Procuração"}).json()["id"]
    assert upload(api, gid, demo_docx).status_code in (200, 201)
    field = api.post(f"/generators/{gid}/fields", json={"label": "OAB do advogado", "type": "oab"}).json()
    assert field["type"] == "oab"
    bad = api.post(f"/generators/{gid}/batch/validate", json={"rows": [{field["name"]: "OAB/XX 1"}]}).json()
    assert "OAB inválida" in str(bad)
    ok = api.post(f"/generators/{gid}/batch/validate", json={"rows": [{field["name"]: "12.345/es"}]}).json()
    assert "OAB inválida" not in str(ok)


def test_suggestion_upload_accept_generate(api, monkeypatch: pytest.MonkeyPatch):
    import io

    from docx import Document

    monkeypatch.setattr("app.modules.suggestions.service.get_ai_provider", lambda: None)
    api.register()
    gid = api.post("/generators", json={"name": "Substabelecimento"}).json()["id"]
    doc = Document()
    doc.add_paragraph("SUBSTABELECIMENTO")
    doc.add_paragraph(
        "OUTORGADO: Ana Lima, advogada, inscrita como OAB/SP 123.456, com escritório em São Paulo."
    )
    buf = io.BytesIO()
    doc.save(buf)
    assert upload(api, gid, buf.getvalue()).status_code in (200, 201)
    items = api.get(f"/generators/{gid}/suggestions").json()["items"]
    (sug,) = [s for s in items if s["suggested_type"] == "oab"]
    field = api.post(f"/suggestions/{sug['id']}/accept").json()["field"]
    assert field["type"] == "oab" and field["name"] == "oab_outorgado"
    job = api.post(
        f"/generators/{gid}/generate", json={"values": {field["name"]: "es 9876-e"}, "format": "docx"}
    ).json()
    (file,) = job["files"]
    text = docx_text(api.tc.get(api.post(f"/files/{file['id']}/link").json()["url"]).content)
    assert "inscrita como OAB/ES 9.876-E, com escritório" in text and "123.456" not in text


MIGRATION_URL = os.environ.get("MIGRATION_DATABASE_URL")


@pytest.mark.skipif(
    not MIGRATION_URL,
    reason="defina MIGRATION_DATABASE_URL (PostgreSQL descartável: o teste recria o schema) para testar a migration",
)
def test_migration_0005_upgrade_downgrade_on_postgres(monkeypatch: pytest.MonkeyPatch):
    from pathlib import Path

    from alembic import command
    from alembic.config import Config

    from app.core.config import reset_settings_cache

    monkeypatch.setenv("DATABASE_URL", MIGRATION_URL)
    reset_settings_cache()
    root = Path(__file__).resolve().parents[1]
    cfg = Config(str(root / "alembic.ini"))
    cfg.set_main_option("script_location", str(root / "migrations"))
    engine = sa.create_engine(MIGRATION_URL.replace("postgresql://", "postgresql+psycopg://", 1))

    def check_def(table: str) -> str:
        with engine.connect() as c:
            return c.execute(
                sa.text("select pg_get_constraintdef(oid) from pg_constraint where conname = :n"),
                {"n": f"ck_{table}_type"},
            ).scalar_one()

    # banco descartável dedicado ao teste (MIGRATION_DATABASE_URL): começa sempre do zero
    with engine.begin() as c:
        c.execute(sa.text("drop schema public cascade"))
        c.execute(sa.text("create schema public"))
    command.upgrade(cfg, "0004")
    assert "oab" not in check_def("fields") and "oab" not in check_def("kit_fields")
    command.upgrade(cfg, "head")
    assert "oab" in check_def("fields") and "oab" in check_def("kit_fields")

    # um campo "oab" existente: downgrade converte para texto curto em vez de falhar
    with engine.begin() as c:
        ws = c.execute(
            sa.text(
                "insert into workspaces (id, name, slug, created_at, updated_at) values (gen_random_uuid(), 'W', 'w-mig', now(), now()) returning id"
            )
        ).scalar_one()
        gen = c.execute(
            sa.text(
                "insert into document_generators (id, workspace_id, name, status, is_pinned, filename_pattern, created_at, updated_at) values (gen_random_uuid(), :w, 'G', 'active', false, '{modelo}', now(), now()) returning id"
            ),
            {"w": ws},
        ).scalar_one()
        ver = c.execute(
            sa.text(
                "insert into template_versions (id, workspace_id, generator_id, version_number, revision, analysis_status, created_at, updated_at) values (gen_random_uuid(), :w, :g, 1, 1, 'done', now(), now()) returning id"
            ),
            {"w": ws, "g": gen},
        ).scalar_one()
        c.execute(
            sa.text(
                "insert into fields (id, workspace_id, version_id, name, label, type, required, include_in_batch, position, created_at, updated_at) values (gen_random_uuid(), :w, :v, 'oab', 'OAB', 'oab', false, true, 0, now(), now())"
            ),
            {"w": ws, "v": ver},
        )
    command.downgrade(cfg, "0004")
    with engine.connect() as c:
        assert c.execute(sa.text("select type from fields where name = 'oab'")).scalar_one() == "text_short"
    assert "oab" not in check_def("fields")
    command.upgrade(cfg, "head")
    command.check(cfg)  # modelos e banco sem divergência depois do ciclo completo
    engine.dispose()
    reset_settings_cache()
