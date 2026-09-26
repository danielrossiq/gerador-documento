"""Auxiliares compartilhados pelos testes do motor DOCX."""

from __future__ import annotations

import io
import zipfile
from collections.abc import Mapping, Sequence

from docx import Document

from app.docx_engine.compile import BlockSpec, CompileResult, compile_docx
from app.docx_engine.nodes import replace_range_with_field
from app.docx_engine.parse import ParsedDocument, parse_docx
from app.docx_engine.render import CompiledTemplate


def specs_from_parsed(
    parsed: ParsedDocument,
    fields: Mapping[str, Sequence[tuple[int, int, str]]] | None = None,
    overrides: Mapping[str, list[dict]] | None = None,
) -> list[BlockSpec]:
    """Converte blocos lidos em BlockSpec, aplicando campos (start, end, field_id) e conteúdos alterados."""
    fields = fields or {}
    overrides = overrides or {}
    specs = []
    for b in parsed.blocks:
        content = b.content
        if b.key in overrides:
            content = overrides[b.key]
        if b.kind == "paragraph":
            for start, end, fid in sorted(fields.get(b.key, ()), reverse=True):
                content, _ = replace_range_with_field(content, start, end, fid)
        specs.append(BlockSpec(b.key, b.container_key, b.ordinal, b.kind, "original", content, None))
    return specs


def compile_with(
    original: bytes,
    fields=None,
    overrides=None,
    names: Mapping[str, str] | None = None,
    extra_specs: Sequence[BlockSpec] = (),
) -> tuple[ParsedDocument, CompileResult]:
    parsed = parse_docx(original)
    specs = specs_from_parsed(parsed, fields, overrides) + list(extra_specs)
    return parsed, compile_docx(original, specs, names or {}, nonce="deadbeef")


def render(compiled: bytes, values: Mapping[str, str]) -> bytes:
    return CompiledTemplate(compiled).render(values)


def all_text(docx_bytes: bytes) -> dict[str, list[str]]:
    """Texto por região: body, tables, header, footer."""
    doc = Document(io.BytesIO(docx_bytes))
    out = {
        "body": [p.text for p in doc.paragraphs],
        "tables": [c.text for t in doc.tables for row in t.rows for c in row.cells],
        "header": [p.text for s in doc.sections for p in s.header.paragraphs],
        "footer": [p.text for s in doc.sections for p in s.footer.paragraphs],
    }
    return out


def flat_text(docx_bytes: bytes) -> str:
    return "\n".join(t for lst in all_text(docx_bytes).values() for t in lst)


def part_xml(docx_bytes: bytes, name: str = "word/document.xml") -> str:
    return zipfile.ZipFile(io.BytesIO(docx_bytes)).read(name).decode("utf-8")


def offset_of(parsed: ParsedDocument, key: str, needle: str, occurrence: int = 0) -> tuple[int, int]:
    block = next(b for b in parsed.blocks if b.key == key)
    idx = -1
    for _ in range(occurrence + 1):
        idx = block.text.index(needle, idx + 1)
    return idx, idx + len(needle)
