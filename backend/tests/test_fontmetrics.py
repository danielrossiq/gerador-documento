"""Compensação de largura de fontes ausentes no worker (docx_engine.fontmetrics)."""

from __future__ import annotations

import io
import zipfile

from docx import Document

from app.docx_engine import fontmetrics
from app.docx_engine.pdf_import import pdf_to_docx
from tests import pdf_samples

INSTALLED = {"liberation sans", "liberation serif", "liberation mono", "carlito", "caladea", "dejavu sans"}


def _docx(font: str, runs: list[tuple[str, str | None]]) -> bytes:
    doc = Document()
    doc.styles["Normal"].font.name = font
    p = doc.add_paragraph()
    for text, run_font in runs:
        r = p.add_run(text)
        if run_font:
            r.font.name = run_font
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _scales(docx: bytes) -> list[str | None]:
    xml = zipfile.ZipFile(io.BytesIO(docx)).read("word/document.xml").decode()
    import re

    out = []
    for run in re.findall(r"<w:r>.*?</w:r>", xml, flags=re.S):
        m = re.search(r'<w:w w:val="(\d+)"/>', run)
        out.append(m.group(1) if m else None)
    return out


def test_missing_font_gets_the_measured_width_and_installed_fonts_are_untouched():
    src = _docx(
        "Consolas", [("Eu, MARIA, declaro", None), (" em Arial", "Arial"), (" em Georgia", "Georgia")]
    )
    out = fontmetrics.compensate_widths(src, INSTALLED)
    assert _scales(out) == [
        "92",
        None,
        "109",
    ]  # Consolas 8% mais estreita; Arial tem métrica igual; Georgia +9%
    assert _scales(fontmetrics.compensate_widths(src, INSTALLED | {"consolas", "georgia"})) == [
        None,
        None,
        None,
    ]


def test_compensation_never_breaks_the_pdf():
    assert fontmetrics.compensate_widths(b"isto nao e um docx", INSTALLED) == b"isto nao e um docx"


def test_pdf_conversion_records_the_ratio_of_its_fonts():
    """Courier do PDF de exemplo tem a mesma largura da Liberation Mono: nada a compensar; uma fonte mais estreita
    teria a proporção gravada em w:docVars, lida pelo worker."""
    conv = pdf_to_docx(pdf_samples.ruled_declaration_pdf())
    assert fontmetrics.read_scale_var(zipfile.ZipFile(io.BytesIO(conv.docx))) == {}
    tagged = fontmetrics.write_scale_var(conv.docx, {"Fonte Rara": 0.88})
    zf = zipfile.ZipFile(io.BytesIO(tagged))
    assert fontmetrics.read_scale_var(zf) == {"fonte rara": 0.88}
    settings = zf.read("word/settings.xml").decode()
    assert settings.index("<w:docVars>") < settings.index("<w:rsids>") if "<w:rsids>" in settings else True


def test_ratio_needs_enough_sample_and_a_real_difference():
    assert fontmetrics.ratio_from_samples(91.7, 100, 200) == 0.917
    assert fontmetrics.ratio_from_samples(99.5, 100, 200) is None  # diferença irrelevante
    assert fontmetrics.ratio_from_samples(80, 100, 10) is None  # amostra pequena
