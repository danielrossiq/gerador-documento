"""Editor (salvamento reconciliado), versões (histórico, restauração), re-ancoragem e sugestões."""

from __future__ import annotations

import copy

from app.docx_engine.inline import inline_text
from tests.flows import (
    SAMPLE_VALUES,
    apply_field,
    build_generator,
    create_field,
    docx_text,
    iter_paragraphs,
    upload,
)


def test_stale_revision_is_rejected(api, demo_docx):
    api.register()
    gid = api.post("/generators", json={"name": "G"}).json()["id"]
    upload(api, gid, demo_docx)
    ed = api.get(f"/generators/{gid}/editor").json()
    ok = api.put(
        f"/generators/{gid}/editor/content", json={"revision": ed["version"]["revision"], "doc": ed["doc"]}
    )
    assert ok.status_code == 200
    stale = api.put(
        f"/generators/{gid}/editor/content", json={"revision": ed["version"]["revision"], "doc": ed["doc"]}
    )
    assert stale.status_code == 409 and stale.json()["error"]["code"] == "STALE_REVISION"
    assert stale.json()["error"]["details"]["current_revision"] == ok.json()["revision"]


def test_editor_doc_structure_has_parts_tables_and_atoms(api, demo_docx):
    api.register()
    gid = api.post("/generators", json={"name": "G"}).json()["id"]
    upload(api, gid, demo_docx)
    ed = api.get(f"/generators/{gid}/editor").json()
    parts = [p["attrs"]["part"] for p in ed["doc"]["content"]]
    assert parts == ["header1", "body", "footer1"]
    body = ed["doc"]["content"][1]["content"]
    assert any(n["type"] == "table" for n in body)
    assert any(
        c["type"] == "docAtom" for n in body if n["type"] == "paragraph" for c in n.get("content", [])
    )  # imagem/quebra
    assert ed["stats"]["images"] == 1 and any("avançado" in w for w in ed["warnings"])
    cell = next(n for n in body if n["type"] == "table")["content"][1]["content"][1]
    assert (
        cell["attrs"]["container"] == "body/t1/r2/c2"
        and cell["content"][0]["attrs"]["blockKey"] == "body/t1/r2/c2/p1"
    )


def test_editor_save_handles_new_and_deleted_paragraphs_and_compiles(api, demo_docx):
    api.register()
    built = build_generator(api, demo_docx)
    gid = built["id"]
    ed = api.get(f"/generators/{gid}/editor").json()
    doc = copy.deepcopy(ed["doc"])
    body = doc["content"][1]["content"]
    idx = next(
        i for i, n in enumerate(body) if n["type"] == "paragraph" and n["attrs"]["blockKey"] == "body/p10"
    )
    body.pop(idx)  # apaga a linha de assinatura
    body.insert(
        idx,
        {
            "type": "paragraph",
            "attrs": {"blockKey": None},
            "content": [{"type": "text", "text": "Parágrafo criado pelo usuário."}],
        },
    )
    saved = api.put(
        f"/generators/{gid}/editor/content", json={"revision": ed["version"]["revision"], "doc": doc}
    )
    assert saved.status_code == 200 and saved.json()["deleted_blocks"] == 1
    # sem etapa de publicação: a alteração já vale para a próxima geração
    job = api.post(f"/generators/{gid}/generate", json={"values": SAMPLE_VALUES}).json()
    url = api.post(f"/files/{job['files'][0]['id']}/link").json()["url"]
    text = docx_text(api.tc.get(url).content)
    assert "Parágrafo criado pelo usuário." in text and "_______________________________" not in text


def test_table_structure_cannot_be_changed(api, demo_docx):
    api.register()
    gid = api.post("/generators", json={"name": "G"}).json()["id"]
    upload(api, gid, demo_docx)
    ed = api.get(f"/generators/{gid}/editor").json()
    doc = copy.deepcopy(ed["doc"])
    body = doc["content"][1]["content"]
    body[:] = [n for n in body if n["type"] != "table"]  # remover a tabela inteira
    r = api.put(f"/generators/{gid}/editor/content", json={"revision": ed["version"]["revision"], "doc": doc})
    assert r.status_code == 422 and r.json()["error"]["code"] == "STRUCTURE_LOCKED"


def test_unknown_field_node_is_healed_to_plain_text(api, demo_docx):
    api.register()
    gid = api.post("/generators", json={"name": "G"}).json()["id"]
    upload(api, gid, demo_docx)
    ed = api.get(f"/generators/{gid}/editor").json()
    doc = copy.deepcopy(ed["doc"])
    p = next(p for p in iter_paragraphs(doc) if p["attrs"]["blockKey"] == "body/p3")
    p["content"].append(
        {"type": "field", "attrs": {"fieldId": "00000000-0000-0000-0000-000000000000", "sample": "fantasma"}}
    )
    r = api.put(f"/generators/{gid}/editor/content", json={"revision": ed["version"]["revision"], "doc": doc})
    assert r.status_code == 200 and any("não existe mais" in w for w in r.json()["warnings"])
    again = api.get(f"/generators/{gid}/editor").json()
    p2 = next(p for p in iter_paragraphs(again["doc"]) if p["attrs"]["blockKey"] == "body/p3")
    assert inline_text(p2["content"]).endswith("fantasma") and all(
        n["type"] != "field" for n in p2["content"]
    )


def test_editing_is_immediate_no_publish_step(api, demo_docx):
    """Não existe mais rascunho/publicação: salvar já deixa a alteração pronta para a próxima geração."""
    api.register()
    built = build_generator(api, demo_docx)
    gid = built["id"]
    v1 = api.get(f"/generators/{gid}/versions").json()["items"]
    assert [v["version_number"] for v in v1] == [1]

    ed = api.get(f"/generators/{gid}/editor").json()
    assert ed["version"]["version_number"] == 1
    assert len(ed["fields"]) == 7 and all(f["occurrences"] > 0 for f in ed["fields"])

    job1 = api.post(f"/generators/{gid}/generate", json={"values": SAMPLE_VALUES}).json()
    url1 = api.post(f"/files/{job1['files'][0]['id']}/link").json()["url"]
    assert "São Paulo, 21 de setembro de 2026." in docx_text(api.tc.get(url1).content)

    doc = copy.deepcopy(ed["doc"])
    p = next(p for p in iter_paragraphs(doc) if p["attrs"]["blockKey"] == "body/p9")
    p["content"] = [{"type": "text", "text": "Cidade nova, 1 de janeiro de 2030."}]
    api.put(f"/generators/{gid}/editor/content", json={"revision": ed["version"]["revision"], "doc": doc})

    job2 = api.post(f"/generators/{gid}/generate", json={"values": SAMPLE_VALUES}).json()  # sem publicar nada
    assert job2["version_number"] == 1  # ainda a mesma versão: só o conteúdo dela mudou
    url2 = api.post(f"/files/{job2['files'][0]['id']}/link").json()["url"]
    assert "Cidade nova, 1 de janeiro de 2030." in docx_text(api.tc.get(url2).content)

    # o primeiro job já gerado não é afetado pela edição (o compilado foi congelado na criação do job)
    assert "São Paulo, 21 de setembro de 2026." in docx_text(api.tc.get(url1).content)

    # campos continuam editáveis livremente (não há "versão publicada" travada)
    field = built["fields"]["nome"]
    ok = api.patch(f"/fields/{field['id']}", json={"label": "Contratante"})
    assert ok.status_code == 200 and ok.json()["label"] == "Contratante"


def test_reupload_creates_new_current_version_without_affecting_past_jobs(api, demo_docx):
    api.register()
    built = build_generator(api, demo_docx)
    gid = built["id"]
    job1 = api.post(f"/generators/{gid}/generate", json={"values": SAMPLE_VALUES}).json()
    assert job1["version_number"] == 1

    upload(api, gid, demo_docx, "v2.docx")  # novo arquivo: vira a versão atual na hora
    ed = api.get(f"/generators/{gid}/editor").json()
    assert ed["version"]["version_number"] == 2
    hist = [v["version_number"] for v in api.get(f"/generators/{gid}/versions").json()["items"]]
    assert hist == [2, 1]
    # a interface mostra qual arquivo está em uso (etapa "Modelo") e o de cada versão do histórico
    assert ed["version"]["source_filename"] == "v2.docx"
    names = [v["source_filename"] for v in api.get(f"/generators/{gid}/versions").json()["items"]]
    assert names[0] == "v2.docx" and names[1].endswith(".docx")

    job2 = api.post(f"/generators/{gid}/generate", json={"values": SAMPLE_VALUES}).json()
    assert job2["version_number"] == 2
    # o job antigo continua apontando para a v1 e seus arquivos continuam baixáveis
    assert api.get(f"/jobs/{job1['id']}").json()["version_number"] == 1
    assert api.post(f"/files/{job1['files'][0]['id']}/link").status_code == 200


def test_restore_old_version_becomes_current(api, demo_docx):
    api.register()
    built = build_generator(api, demo_docx)
    gid = built["id"]
    upload(api, gid, demo_docx, "v2.docx")  # cria a v2 (atual)
    old = api.get(f"/generators/{gid}/versions").json()["items"][-1]  # v1
    assert old["version_number"] == 1

    restored = api.post(f"/generators/{gid}/versions/{old['id']}/restore")
    assert (
        restored.status_code == 201 and restored.json()["version_number"] == 3
    )  # nunca reaproveita o número
    assert api.get(f"/generators/{gid}/editor").json()["version"]["version_number"] == 3

    # restaurar a própria versão atual é um conflito claro, não um no-op silencioso
    current_id = api.get(f"/generators/{gid}/versions").json()["items"][0]["id"]
    again = api.post(f"/generators/{gid}/versions/{current_id}/restore")
    assert again.status_code == 409 and again.json()["error"]["code"] == "ALREADY_CURRENT"


def test_reupload_reanchors_fields_and_flags_missing_ones(api, demo_docx):
    """Novo arquivo: as definições de campo são mantidas; ocorrências são re-ancoradas por trecho+contexto."""
    import io

    from docx import Document

    from app.modules.samples.documents import build_demo_contract

    api.register()
    built = build_generator(api, demo_docx)
    gid = built["id"]
    # arquivo novo SEM o valor "R$ 5.000,00" (troca por texto fixo)
    d = Document(io.BytesIO(build_demo_contract()))
    for p in d.paragraphs:
        for r in p.runs:
            if r.text == "R$ 5.000,00":
                r.text = "valor a combinar"
    for row in d.tables[0].rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    r.text = r.text.replace("R$ 5.000,00", "a combinar")
    buf = io.BytesIO()
    d.save(buf)
    up = upload(api, gid, buf.getvalue(), "v2.docx")
    assert up.status_code == 201
    fields = {f["name"]: f for f in api.get(f"/generators/{gid}/editor").json()["fields"]}
    assert fields["nome_do_contratante"]["occurrences"] == 5  # re-ancorado
    assert fields["cpf_do_contratante"]["occurrences"] == 1
    assert (
        fields["valor_mensal"]["occurrences"] == 0 and fields["valor_mensal"]["missing_in_document"] is True
    )


def test_generation_needs_a_document_first(api, demo_docx):
    """Sem publicar nada: um gerador sem arquivo enviado ainda não pode gerar; assim que tem campos, já pode."""
    api.register()
    empty_gid = api.post("/generators", json={"name": "Sem modelo"}).json()["id"]
    empty = api.post(f"/generators/{empty_gid}/generate", json={"values": {}})
    assert empty.status_code == 404 and empty.json()["error"]["code"] == "NOT_FOUND"

    gid = build_generator(api, demo_docx)["id"]
    r = api.post(f"/generators/{gid}/generate", json={"values": SAMPLE_VALUES})
    assert r.status_code == 202 and r.json()["status"] == "succeeded"
    assert api.get("/dashboard").json()["documents_this_month"] == 1  # toda geração conta (sem modo de teste)


def test_suggestions_review_accept_ignore_and_provider_label(api, demo_docx):
    api.register()
    gid = api.post("/generators", json={"name": "G"}).json()["id"]
    upload(api, gid, demo_docx)  # análise roda no upload (backend 'sync')
    data = api.get(f"/generators/{gid}/suggestions").json()
    assert (
        data["analysis"]["status"] == "done"
        and "IA externa não configurada" in data["analysis"]["provider_label"]
    )
    by_name = {g["suggested_name"]: g for g in data["items"]}
    assert by_name["cpf_contratante"]["confidence"] >= 0.95 and by_name["nome_contratante"]["count"] == 5
    assert by_name["valor"]["count"] == 2 and by_name["valor"]["suggested_type"] == "currency"

    acc = api.post(
        f"/suggestions/{by_name['nome_contratante']['id']}/accept", json={"label": "Nome do contratante"}
    )
    assert acc.status_code == 200
    out = acc.json()
    assert (
        out["applied"] == 5
        and out["field"]["occurrences"] == 5
        and out["field"]["name"] == "nome_contratante"
    )
    # aceitar sugestão de lista suspensa já cria a primeira opção
    sel = api.post(f"/suggestions/{by_name['tipo_servico']['id']}/accept").json()
    assert sel["field"]["type"] == "select" and [o["label"] for o in sel["field"]["options"]] == [
        "consultoria em desenvolvimento de software"
    ]
    ign = api.post(f"/suggestions/{by_name['valor']['id']}/ignore")
    assert ign.json()["ignored"] == 2
    left = {g["suggested_name"] for g in api.get(f"/generators/{gid}/suggestions").json()["items"]}
    assert "nome_contratante" not in left and "valor" not in left and "tipo_servico" not in left
    # reanálise não ressuscita o que foi ignorado
    api.post(f"/generators/{gid}/suggestions/analyze", json={"force": True})
    left2 = {g["suggested_name"] for g in api.get(f"/generators/{gid}/suggestions").json()["items"]}
    assert "valor" not in left2
    # o documento agora contém os campos aceitos
    ed = api.get(f"/generators/{gid}/editor").json()
    assert any(n["type"] == "field" for p in iter_paragraphs(ed["doc"]) for n in p.get("content", []))


def test_accept_all_high_confidence(api, demo_docx):
    api.register()
    gid = api.post("/generators", json={"name": "G"}).json()["id"]
    upload(api, gid, demo_docx)
    r = api.post(f"/generators/{gid}/suggestions/accept-all", json={"min_confidence": 0.9}).json()
    assert r["accepted_groups"] >= 4
    names = {f["name"] for f in api.get(f"/generators/{gid}/fields").json()["items"]}
    assert {"cpf_contratante", "cnpj_contratada", "valor"} <= names


def test_stale_suggestion_disappears_and_never_creates_an_empty_field(api, demo_docx):
    """O usuário marca "João da Silva" à mão e depois clica em "Aceitar todas": a sugestão de nome do
    contratante (que já virou campo) não pode criar um segundo campo sem ocorrência nenhuma no documento."""
    import uuid

    from sqlalchemy import select

    from app.db.models import FieldSuggestion
    from app.db.session import get_factory

    api.register()
    gid = build_generator(api, demo_docx)["id"]  # já transformou nome, CPF, valor… em campos no editor
    left = {g["suggested_name"] for g in api.get(f"/generators/{gid}/suggestions").json()["items"]}
    assert "nome_contratante" not in left and "cpf_contratante" not in left
    assert "nome_contratada" in left  # "Estúdio Exemplo Serviços Ltda." continua sendo texto fixo

    version_id = uuid.UUID(api.get(f"/generators/{gid}/editor").json()["version"]["id"])
    with get_factory()() as db:
        stale_id = db.scalar(
            select(FieldSuggestion.id).where(
                FieldSuggestion.version_id == version_id,
                FieldSuggestion.suggested_name == "nome_contratante",
                FieldSuggestion.status == "pending",
            )
        )
    assert stale_id is not None
    before = {f["name"] for f in api.get(f"/generators/{gid}/fields").json()["items"]}
    r = api.post(f"/suggestions/{uuid.UUID(str(stale_id))}/accept")
    assert r.status_code == 409 and r.json()["error"]["code"] == "SUGGESTION_STALE"
    assert {f["name"] for f in api.get(f"/generators/{gid}/fields").json()["items"]} == before

    api.post(f"/generators/{gid}/suggestions/accept-all", json={"min_confidence": 0.6})
    fields = api.get(f"/generators/{gid}/editor").json()["fields"]
    assert all(f["occurrences"] > 0 for f in fields)
    assert "nome_contratada" in {f["name"] for f in fields}


def test_accept_month_suggestion_seeds_all_twelve_options_not_just_the_matched_text(api, demo_docx):
    """Ao contrário de um "select" comum (que só recebe a primeira opção vista, ver o teste de
    tipo_servico acima), a sugestão de "mês" já vem com a lista inteira pronta (`config.options`,
    ver app.modules.suggestions.rules.MONTH_OPTIONS) — aceitar cria o campo com os 12 meses, e o
    "options" não deve vazar para o `config` do campo criado."""
    import uuid

    from app.db.models import FieldSuggestion, TemplateVersion
    from app.db.session import get_factory
    from app.modules.fields.normalize import MONTHS
    from app.modules.suggestions.rules import MONTH_OPTIONS

    api.register()
    gid = api.post("/generators", json={"name": "G"}).json()["id"]
    upload(api, gid, demo_docx)
    ed = api.get(f"/generators/{gid}/editor").json()
    version_id = ed["version"]["id"]
    # o contrato de exemplo tem "21 de setembro de 2026" na assinatura: a sugestão precisa apontar para
    # um trecho que existe (sugestão sem nenhuma ocorrência é descartada, ver test_stale_suggestion_*)
    block_key, start = next(
        (p["attrs"]["blockKey"], inline_text(p.get("content", [])).index("setembro"))
        for p in iter_paragraphs(ed["doc"])
        if "setembro" in inline_text(p.get("content", []))
    )

    with get_factory()() as db:
        version = db.get(TemplateVersion, uuid.UUID(version_id))
        sug = FieldSuggestion(
            workspace_id=version.workspace_id,
            version_id=version.id,
            block_key=block_key,
            start_offset=start,
            end_offset=start + len("setembro"),
            text="setembro",
            context_before="",
            context_after="",
            suggested_name="mes",
            suggested_label="Mês",
            suggested_type="select",
            suggested_config={"options": MONTH_OPTIONS},
            confidence=0.6,
            source="rule",
            reason="Nome de mês avulso no texto",
            location="Parágrafo 1",
            status="pending",
        )
        db.add(sug)
        db.commit()
        sug_id = str(sug.id)

    acc = api.post(f"/suggestions/{sug_id}/accept")
    assert acc.status_code == 200, acc.text
    field = acc.json()["field"]
    assert field["type"] == "select"
    assert [o["label"] for o in field["options"]] == [m.capitalize() for m in MONTHS]
    assert not field["config"]  # "options" não vaza pro config do campo (API serializa ausência como {})


def test_dashboard_and_list_filters(api, demo_docx):
    api.register()
    build_generator(api, demo_docx, name="Contrato A")
    b = api.post("/generators", json={"name": "Rascunho B", "description": "sem modelo"}).json()
    api.post("/generators", json={"name": "Termo C"})
    lst = api.get("/generators?q=contrato").json()
    assert (
        lst["total"] == 1 and lst["items"][0]["name"] == "Contrato A" and lst["items"][0]["fields_count"] == 7
    )
    assert api.get("/generators").json()["total"] == 3  # sem publicar nada, os 3 aparecem normalmente
    assert api.get("/generators?sort=name").json()["items"][0]["name"] == "Contrato A"
    dash = api.get("/dashboard").json()
    assert dash["total_generators"] == 3 and len(dash["recent"]) == 3

    # fixar traz o gerador para o topo da listagem, à frente da ordenação escolhida
    assert api.post(f"/generators/{b['id']}/pin").json()["is_pinned"] is True
    assert api.get("/generators?sort=name").json()["items"][0]["name"] == "Rascunho B"
    assert api.post(f"/generators/{b['id']}/unpin").json()["is_pinned"] is False
    assert api.get("/generators?sort=name").json()["items"][0]["name"] == "Contrato A"


def test_duplicate_archive_delete_generator(api, demo_docx):
    api.register()
    gid = build_generator(api, demo_docx)["id"]
    dup = api.post(f"/generators/{gid}/duplicate")
    assert (
        dup.status_code == 201 and dup.json()["name"].endswith("(cópia)") and dup.json()["fields_count"] == 7
    )
    did = dup.json()["id"]
    ed = api.get(f"/generators/{did}/editor").json()
    assert len(ed["fields"]) == 7 and all(f["occurrences"] > 0 for f in ed["fields"])
    assert api.post(f"/generators/{gid}/archive").json()["status"] == "archived"
    assert api.get("/generators").json()["total"] == 1
    assert api.get("/generators?status=archived").json()["total"] == 1
    assert api.post(f"/generators/{gid}/unarchive").json()["status"] == "active"
    assert api.delete(f"/generators/{did}").status_code == 204
    assert api.get(f"/generators/{did}").status_code == 404  # exclusão lógica: some do usuário
    assert api.get("/generators").json()["total"] == 1


def test_field_reorder_and_patch(api, demo_docx):
    api.register()
    built = build_generator(api, demo_docx)
    gid = built["id"]
    ids = [f["id"] for f in api.get(f"/generators/{gid}/fields").json()["items"]]
    assert api.put(f"/generators/{gid}/fields/order", json={"ids": list(reversed(ids))}).status_code == 204
    assert [f["id"] for f in api.get(f"/generators/{gid}/fields").json()["items"]] == list(reversed(ids))
    assert api.put(f"/generators/{gid}/fields/order", json={"ids": ids[:-1]}).status_code == 422
    nome = built["fields"]["nome"]
    r = api.patch(
        f"/fields/{nome['id']}",
        json={"label": "Contratante", "required": False, "description": "Quem contrata"},
    )
    assert r.status_code == 200 and r.json()["label"] == "Contratante" and r.json()["occurrences"] == 5
    bad = api.patch(f"/fields/{nome['id']}", json={"name": "Nome Inválido"})
    assert bad.status_code == 422 and bad.json()["error"]["code"] == "INVALID_FIELD_NAME"
    res = api.patch(f"/fields/{nome['id']}", json={"name": "linha"})
    assert res.json()["error"]["code"] == "RESERVED_FIELD_NAME"
    _ = create_field, apply_field


def test_queue_outage_does_not_break_upload_or_leave_job_queued(api, demo_docx, monkeypatch):
    """Redis/fila fora do ar: o upload funciona (análise marcada como falha) e a geração falha com erro claro."""
    from tests.flows import SAMPLE_VALUES, build_generator

    api.register()
    boom = RuntimeError("redis indisponível")

    def broken(*a, **k):
        raise boom

    monkeypatch.setattr("app.modules.templates.service.enqueue", broken)
    gid = api.post("/generators", json={"name": "G"}).json()["id"]
    up = upload(api, gid, demo_docx)
    assert up.status_code == 201
    assert (
        up.json()["version"]["analysis_status"] == "failed"
        or api.get(f"/generators/{gid}/editor").json()["version"]["analysis_error"]
    )
    monkeypatch.undo()
    built = build_generator(api, demo_docx)
    monkeypatch.setattr("app.modules.generation.service.enqueue", broken)
    r = api.post(f"/generators/{built['id']}/generate", json={"values": SAMPLE_VALUES})
    assert r.status_code == 503 and r.json()["error"]["code"] == "QUEUE_UNAVAILABLE"
    jobs = api.get("/jobs").json()["items"]
    assert jobs and jobs[0]["status"] == "failed" and jobs[0]["error_code"] == "QUEUE_UNAVAILABLE"
